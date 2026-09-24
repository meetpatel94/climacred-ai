#!/usr/bin/env python3
"""
ClimaCred AI - Test Data Pack builder (TESTING ARTEFACT, NOT APPLICATION CODE)
=============================================================================

This script generates the ClimaCred_AI_Test_Data_Pack. It is a *separate testing
artefact*: nothing in the ClimaCred application imports it, and it does not
modify application behaviour. It only READS the repository's own climate engine
(``backend/app/climate_engine``) so that every score in the pack is produced by
the exact code the application runs.

Design rule: ONE monthly model per business. Every table (energy, water, waste,
mobility, emissions, materials, history, before/after) is derived from that same
model, so the files cannot drift apart.

All businesses, people and numbers are fictional demo data.
"""
from __future__ import annotations

import csv
import json
import math
import os
import random
import sys
import zipfile
from datetime import datetime, timezone

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
BACKEND = os.path.join(REPO_ROOT, "backend")
OUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))

sys.path.insert(0, BACKEND)

# --- the application's own engine (read-only import) -------------------------
from app.config import settings  # noqa: E402
from app.climate_engine.fingerprint import generate_fingerprint, score_label_from_overall  # noqa: E402
from app.climate_engine.recommendations import SOLUTION_CATALOG  # noqa: E402
from app.utils.calculations import calculate_emissions_metrics  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

# ---------------------------------------------------------------------------
# Factors. Grid / diesel / petrol / gas factors are read from the application
# config so the pack can never disagree with the app. LPG, CNG and EV charging
# reuse the app's gaseous / electricity factors (documented in the README).
# ---------------------------------------------------------------------------
EF_ELEC = settings.ELECTRICITY_EMISSION_FACTOR_KG_PER_KWH      # 0.82 kg CO2e/kWh
EF_DIESEL = settings.DIESEL_EMISSION_FACTOR_KG_PER_LITRE       # 2.68
EF_PETROL = settings.PETROL_EMISSION_FACTOR_KG_PER_LITRE       # 2.31
EF_GAS = settings.NATURAL_GAS_EMISSION_FACTOR_KG_PER_KG        # 2.75 (also used for LPG/CNG)
WATER_COST_PER_KL = settings.WATER_COST_PER_KL_INR             # 30

# Scope 3 factors (illustrative, documented, identical for every business)
S3_MATERIAL_FACTOR = {   # kg CO2e per kg purchased material
    "cotton_yarn": 5.90, "steel": 2.10, "agri_produce": 0.55,
    "kraft_paper": 1.05, "polymer": 2.40, "wood": 0.46,
}
S3_FREIGHT_KG_PER_KM = 0.42        # third-party heavy goods vehicle
S3_LANDFILL_KG_PER_KG = 0.55       # mixed industrial waste to landfill
S3_TRAVEL_KG_PER_KM = 0.12         # mixed car/rail business travel

MONTHS = []
_y, _m = 2024, 9
for _ in range(24):
    MONTHS.append(f"{_y:04d}-{_m:02d}")
    _m += 1
    if _m == 13:
        _y, _m = _y + 1, 1
LAST_12 = MONTHS[-12:]
CURRENT_MONTH = MONTHS[-1]                 # 2026-08
ASSESSMENT_DATE = f"{CURRENT_MONTH}-31"

# Indian operating-day calendar (festival shutdowns included)
BASE_DAYS = {
    "2024-09": 25, "2024-10": 26, "2024-11": 23, "2024-12": 26,
    "2025-01": 25, "2025-02": 25, "2025-03": 26, "2025-04": 24, "2025-05": 25,
    "2025-06": 25, "2025-07": 26, "2025-08": 25, "2025-09": 25, "2025-10": 26,
    "2025-11": 22, "2025-12": 26,
    "2026-01": 25, "2026-02": 25, "2026-03": 26, "2026-04": 24, "2026-05": 25,
    "2026-06": 25, "2026-07": 26, "2026-08": 25,
}


def r(v, nd=0):
    """Round to nd decimals; return int when nd == 0 (keeps XLSX cells clean)."""
    v = round(float(v), nd) if nd else round(float(v))
    return v if nd else int(v)


def series(base, trend=0.0, seasonal=None, breakpoints=None, noise=0.0, rng=None):
    """One 24-month series: base * trend^i * seasonal * breakpoints * (1 +/- noise)."""
    out = []
    for i, m in enumerate(MONTHS):
        v = base * ((1.0 + trend) ** i)
        if seasonal:
            v *= seasonal.get(m[5:7], 1.0)
        if breakpoints:
            for bm, mult in breakpoints:
                if m >= bm:
                    v *= mult
        if noise and rng is not None:
            v *= 1.0 + rng.uniform(-noise, noise)
        out.append(v)
    return out


def flag_timeline(timeline, month):
    """Value of a qualitative flag in `month` given [(effective_month, value)]."""
    val = None
    for eff, v in sorted(timeline):
        if month >= eff:
            val = v
    return val


# ---------------------------------------------------------------------------
# BUSINESS DEFINITIONS (fictional)
# ---------------------------------------------------------------------------
SEASON_FLAT = {str(i).zfill(2): 1.0 for i in range(1, 13)}

BUSINESSES = [
    dict(
        bid="B001",
        name="Sunrise Textile Industries Pvt. Ltd.",
        industry="Textile Manufacturing",
        sub_industry="Knitting, Dyeing & Garment Finishing",
        city="Tiruppur", state="Tamil Nadu",
        employees=148, facility_area_sqft=42000,
        days=26, hours=16, turnover=345000000,
        ownership="Private Limited", facility_type="Owned Industrial Shed",
        established=2009,
        products="Knitted cotton fabric; dyed & finished garments",
        email="operations.b001@climacred-demo.example",
        created_at="2026-09-01T09:15:00Z",
        business_size="Medium",
        archetype="High-water / high-emissions business with worsening performance",
        # --- monthly drivers (24-month model) ---
        prod_base=68000, prod_unit="kg", prod_label="kg of dyed & finished fabric",
        kwh_daily=10600, kwh_trend=0.0090, kwh_breaks=[("2026-01", 0.82)],
        kwh_seasonal={"01": 0.99, "02": 1.00, "03": 1.03, "04": 1.06, "05": 1.08, "06": 1.05,
                      "07": 1.00, "08": 1.01, "09": 1.04, "10": 1.06, "11": 1.02, "12": 1.03},
        solar_timeline=[("2024-09", 0.0)],
        eff_equip=25.0, tariff=8.60, genset_hours=58,
        water_daily=184000, water_trend=0.0050,
        water_seasonal={"01": 1.02, "02": 1.00, "03": 1.03, "04": 0.98, "05": 0.96, "06": 0.95,
                        "07": 1.00, "08": 1.01, "09": 1.05, "10": 1.06, "11": 1.02, "12": 1.03},
        water_breaks=[("2026-03", 0.94)],
        ww_fraction=0.76, treat_fraction_tl=[("2024-09", 0.95), ("2025-11", 0.40)],
        process_share=0.80, cleaning_share=0.11,
        reuse_fraction=0.0,
        waste=dict(organic=420, plastic=1250, paper=640, metal=260, material=9800,
                   industrial=1900, general=1100, hazardous=2600),
        waste_trend=0.0040,
        dry_recovery_tl=[("2024-09", 0.42), ("2025-06", 0.40), ("2025-11", 0.38), ("2026-02", 0.364)],
        organic_recovery_tl=[("2024-09", 0.0)],
        genset_diesel=1100, process_diesel=0, lpg=320,
        fleet_diesel=1650, fleet_petrol=0, fleet_cng=0, ev_kwh=0,
        vehicles_tl=[("2024-09", 6)], vehicle_fuel="Diesel",
        fleet_mileage=6.5, logistics_km=14500, s3_freight_km=3775,
        commute_km_per_emp_day=14, commute_factor=0.053, commute_mode="Two-Wheelers",
        travel_km=900,
        mats=dict(cotton_yarn=76500), recycled_input=4200, output_kg_per_unit=1.0,
        rejected=3100, packaging=2400,
        peak_load_factor=1.26,
        disposal_rate=3.2, haz_rate=18.0,
        # assessment qualitative flags
        water_source="Groundwater / Borewell",
        leakage_tl=[("2024-09", "Rarely"), ("2025-08", "Monthly")],
        wastewater_tl=[("2024-09", "Full ETP / STP"), ("2025-11", "Primary / Settling")],
        primary_fuel="Electricity Grid", air_control="Basic Scrubber",
        green_tl={
            "ledLighting": [("2024-09", True)], "solarPanels": [("2024-09", False)],
            "rainwaterHarvesting": [("2024-09", False)], "waterRecycling": [("2024-09", False)],
            "wasteSegregation": [("2024-09", True)],
            "energyEfficientMachinery": [("2024-09", False)],
            "evAdoption": [("2024-09", False)], "sustainableMaterials": [("2024-09", False)],
        },
        noise=0.030, seed=101,
        before_after=True,
        interventions=[("2026-01", "LED lighting retrofit + compressed-air leak repair (Phase 1 quick wins)"),
                       ("2026-03", "Smart ultrasonic water leak & flow telemetry")],
    ),
    dict(
        bid="B002",
        name="TechNova Precision Machining LLP",
        industry="Precision Engineering / Machining",
        sub_industry="CNC Machined Automotive Components",
        city="Rajkot", state="Gujarat",
        employees=62, facility_area_sqft=18500,
        days=26, hours=12, turnover=128000000,
        ownership="Limited Liability Partnership", facility_type="Leased Industrial Unit",
        established=2014,
        products="CNC-machined shafts, housings & precision auto components",
        email="works.b002@climacred-demo.example",
        created_at="2026-09-01T11:40:00Z",
        business_size="Small",
        archetype="High-diesel / high-logistics business, mildly worsening",
        prod_base=28000, prod_unit="pieces", prod_label="machined components (pieces)",
        kwh_daily=3690, kwh_trend=0.0040, kwh_breaks=[],
        kwh_seasonal=SEASON_FLAT,
        solar_timeline=[("2024-09", 0.0)],
        eff_equip=40.0, tariff=8.90, genset_hours=26,
        water_daily=1846, water_trend=0.0015, water_seasonal=SEASON_FLAT, water_breaks=[],
        ww_fraction=0.65, treat_fraction_tl=[("2024-09", 0.70)],
        process_share=0.54, cleaning_share=0.25,
        reuse_fraction=0.12,
        waste=dict(organic=210, plastic=480, paper=260, metal=8150, material=0,
                   industrial=420, general=350, hazardous=1450),
        waste_trend=0.0030,
        dry_recovery_tl=[("2024-09", 0.50), ("2025-08", 0.46), ("2026-05", 0.42)],
        organic_recovery_tl=[("2024-09", 0.0)],
        genset_diesel=420, process_diesel=260, lpg=0,
        fleet_diesel=2850, fleet_petrol=240, fleet_cng=0, ev_kwh=0,
        vehicles_tl=[("2024-09", 8), ("2025-06", 9), ("2026-04", 10)], vehicle_fuel="Diesel",
        fleet_mileage=5.5, logistics_km=26000, s3_freight_km=10325,
        commute_km_per_emp_day=18, commute_factor=0.048, commute_mode="Mixed",
        travel_km=1850,
        mats=dict(steel=33600), recycled_input=12400, output_kg_per_unit=1.05,
        rejected=1150, packaging=1800,
        peak_load_factor=1.25,
        disposal_rate=3.2, haz_rate=18.0,
        water_source="Municipal + Borewell",
        leakage_tl=[("2024-09", "Never"), ("2025-09", "Rarely")],
        wastewater_tl=[("2024-09", "Primary / Settling")],
        primary_fuel="Electricity Grid", air_control="Bag Filter / ESP",
        green_tl={
            "ledLighting": [("2024-09", True)], "solarPanels": [("2024-09", False)],
            "rainwaterHarvesting": [("2024-09", False)], "waterRecycling": [("2024-09", False)],
            "wasteSegregation": [("2024-09", True)],
            "energyEfficientMachinery": [("2024-09", False)],
            "evAdoption": [("2024-09", False)], "sustainableMaterials": [("2024-09", False)],
        },
        noise=0.025, seed=202,
        before_after=False, interventions=[],
    ),
    dict(
        bid="B003",
        name="FreshHarvest Food Processing Pvt. Ltd.",
        industry="Food Processing",
        sub_industry="Frozen Vegetables, Pulp & Dairy Desserts",
        city="Anand", state="Gujarat",
        employees=96, facility_area_sqft=26000,
        days=24, hours=14, turnover=214000000,
        ownership="Private Limited", facility_type="Owned Processing Plant",
        established=2011,
        products="IQF vegetables, fruit pulp & puree, dairy-based desserts",
        email="plant.b003@climacred-demo.example",
        created_at="2026-09-02T08:05:00Z",
        business_size="Medium",
        archetype="High organic waste, strong recycling, clearly improving",
        prod_base=96000, prod_unit="kg", prod_label="kg of processed food product",
        kwh_daily=5750, kwh_trend=-0.0035, kwh_breaks=[],
        kwh_seasonal={"01": 0.96, "02": 0.97, "03": 1.02, "04": 1.14, "05": 1.20, "06": 1.22,
                      "07": 1.16, "08": 1.02, "09": 0.98, "10": 1.04, "11": 1.06, "12": 1.05},
        solar_timeline=[("2024-09", 0.0), ("2025-02", 20.0)],
        eff_equip=55.0, tariff=8.20, genset_hours=18,
        water_daily=11875, water_trend=-0.0010,
        water_seasonal={"01": 0.95, "02": 0.96, "03": 1.01, "04": 1.10, "05": 1.16, "06": 1.18,
                        "07": 1.12, "08": 1.00, "09": 0.97, "10": 1.02, "11": 1.03, "12": 1.02},
        water_breaks=[],
        ww_fraction=0.83, treat_fraction_tl=[("2024-09", 1.0)],
        process_share=0.67, cleaning_share=0.24,
        reuse_fraction=0.155,
        waste=dict(organic=6400, plastic=1150, paper=780, metal=120, material=0,
                   industrial=980, general=620, hazardous=180),
        waste_trend=0.0005,
        dry_recovery_tl=[("2024-09", 0.78), ("2026-01", 0.85)],
        organic_recovery_tl=[("2024-09", 0.35), ("2025-06", 0.48), ("2026-01", 0.55)],
        genset_diesel=260, process_diesel=0, lpg=640,
        fleet_diesel=980, fleet_petrol=340, fleet_cng=210, ev_kwh=480,
        vehicles_tl=[("2024-09", 5)], vehicle_fuel="Mixed Fleet",
        fleet_mileage=7.2, logistics_km=11800, s3_freight_km=6200,
        commute_km_per_emp_day=9, commute_factor=0.028, commute_mode="Company Bus",
        travel_km=1250,
        mats=dict(agri_produce=108000), recycled_input=6800, output_kg_per_unit=1.0,
        rejected=7600, packaging=4600,
        peak_load_factor=1.24,
        disposal_rate=3.2, haz_rate=18.0,
        water_source="Municipal",
        leakage_tl=[("2024-09", "Rarely")],
        wastewater_tl=[("2024-09", "Full ETP / STP")],
        primary_fuel="Natural Gas / PNG", air_control="Bag Filter / ESP",
        green_tl={
            "ledLighting": [("2024-09", True)], "solarPanels": [("2024-09", False), ("2025-02", True)],
            "rainwaterHarvesting": [("2024-09", False)],
            "waterRecycling": [("2024-09", False), ("2024-12", True)],
            "wasteSegregation": [("2024-09", True)],
            "energyEfficientMachinery": [("2024-09", True)],
            "evAdoption": [("2024-09", False), ("2025-08", True)],
            "sustainableMaterials": [("2024-09", False)],
        },
        noise=0.035, seed=303,
        before_after=True,
        interventions=[("2024-12", "Ultrafiltration water-recycling loop commissioned"),
                       ("2025-02", "20 kWp rooftop solar"),
                       ("2025-06", "Organic-waste composting line started")],
    ),
    dict(
        bid="B004",
        name="GreenPack Packaging Works",
        industry="Packaging Manufacturing",
        sub_industry="Corrugated Boxes & Flexible Packaging",
        city="Ahmedabad", state="Gujarat",
        employees=120, facility_area_sqft=22000,
        days=25, hours=12, turnover=96000000,
        ownership="Partnership Firm", facility_type="Owned Industrial Unit",
        established=2016,
        products="Corrugated boxes, moulded pulp trays, flexible laminates",
        email="unit.b004@climacred-demo.example",
        created_at="2026-09-02T10:25:00Z",
        business_size="Small",
        archetype="Relatively efficient business, strong recycling, improving",
        prod_base=185000, prod_unit="units", prod_label="packaging units (pieces)",
        kwh_daily=1760, kwh_trend=-0.0050, kwh_breaks=[],
        kwh_seasonal={"01": 0.98, "02": 0.99, "03": 1.02, "04": 1.03, "05": 1.04, "06": 1.02,
                      "07": 1.03, "08": 1.06, "09": 1.08, "10": 1.07, "11": 1.00, "12": 0.99},
        solar_timeline=[("2024-09", 0.0), ("2025-06", 48.0)],
        eff_equip=78.0, tariff=8.40, genset_hours=6,
        water_daily=2480, water_trend=-0.0020, water_seasonal=SEASON_FLAT, water_breaks=[],
        ww_fraction=0.76, treat_fraction_tl=[("2024-09", 1.0)],
        process_share=0.29, cleaning_share=0.42,
        reuse_fraction=0.127,
        waste=dict(organic=210, plastic=1350, paper=2900, metal=90, material=0,
                   industrial=420, general=560, hazardous=60),
        waste_trend=-0.0020,
        dry_recovery_tl=[("2024-09", 0.88), ("2025-09", 0.95)],
        organic_recovery_tl=[("2024-09", 0.40)],
        genset_diesel=150, process_diesel=0, lpg=120,
        fleet_diesel=420, fleet_petrol=260, fleet_cng=0, ev_kwh=620,
        vehicles_tl=[("2024-09", 3)], vehicle_fuel="Mixed Fleet",
        fleet_mileage=8.4, logistics_km=3800, s3_freight_km=1450,
        commute_km_per_emp_day=8, commute_factor=0.028, commute_mode="Company Bus",
        travel_km=620,
        mats=dict(kraft_paper=68000, polymer=18000), recycled_input=52000,
        output_kg_per_unit=0.42, rejected=3900, packaging=4200,
        peak_load_factor=1.26,
        disposal_rate=3.2, haz_rate=18.0,
        water_source="Municipal",
        leakage_tl=[("2024-09", "Never")],
        wastewater_tl=[("2024-09", "Full ETP / STP")],
        primary_fuel="Electricity Grid", air_control="Bag Filter / ESP",
        green_tl={
            "ledLighting": [("2024-09", True)], "solarPanels": [("2024-09", False), ("2025-06", True)],
            "rainwaterHarvesting": [("2024-09", False), ("2025-01", True)],
            "waterRecycling": [("2024-09", True)],
            "wasteSegregation": [("2024-09", True)],
            "energyEfficientMachinery": [("2024-09", True)],
            "evAdoption": [("2024-09", False), ("2025-04", True)],
            "sustainableMaterials": [("2024-09", True)],
        },
        noise=0.028, seed=404,
        before_after=True,
        interventions=[("2025-01", "Rainwater harvesting & recharge well"),
                       ("2025-06", "48 kWp rooftop solar"),
                       ("2026-02", "Packaging light-weighting / recycled board substitution")],
    ),
    dict(
        bid="B005",
        name="BlueStone Furniture Works",
        industry="Furniture Manufacturing",
        sub_industry="Solid Wood & Engineered-Wood Furniture",
        city="Jaipur", state="Rajasthan",
        employees=38, facility_area_sqft=16000,
        days=25, hours=10, turnover=42000000,
        ownership="Proprietorship", facility_type="Owned Workshop",
        established=2007,
        products="Solid-wood furniture, modular cabinets, upholstered seating",
        email="workshop.b005@climacred-demo.example",
        created_at="2026-09-03T09:50:00Z",
        business_size="Micro",
        archetype="High energy per employee, zero renewables, worsening operations",
        prod_base=1450, prod_unit="pieces", prod_label="furniture pieces",
        kwh_daily=2080, kwh_trend=0.0045, kwh_breaks=[],
        kwh_seasonal={"01": 1.00, "02": 1.01, "03": 1.03, "04": 1.04, "05": 1.06, "06": 1.05,
                      "07": 1.00, "08": 1.01, "09": 1.03, "10": 1.08, "11": 1.06, "12": 1.05},
        solar_timeline=[("2024-09", 0.0)],
        eff_equip=22.0, tariff=9.30, genset_hours=46,
        water_daily=1520, water_trend=0.0020, water_seasonal=SEASON_FLAT, water_breaks=[],
        ww_fraction=0.71, treat_fraction_tl=[("2024-09", 0.0)],
        process_share=0.32, cleaning_share=0.23,
        reuse_fraction=0.0,
        waste=dict(organic=90, plastic=280, paper=210, metal=60, material=2650,
                   industrial=1150, general=340, hazardous=420),
        waste_trend=0.0030,
        dry_recovery_tl=[("2024-09", 0.42), ("2025-09", 0.35)],
        organic_recovery_tl=[("2024-09", 0.0)],
        genset_diesel=340, process_diesel=60, lpg=260,
        fleet_diesel=180, fleet_petrol=240, fleet_cng=0, ev_kwh=0,
        vehicles_tl=[("2024-09", 2)], vehicle_fuel="Petrol",
        fleet_mileage=9.6, logistics_km=1450, s3_freight_km=620,
        commute_km_per_emp_day=12, commute_factor=0.053, commute_mode="Two-Wheelers",
        travel_km=380,
        mats=dict(wood=62000), recycled_input=6200, output_kg_per_unit=28.0,
        rejected=4300, packaging=2900,
        peak_load_factor=1.25,
        disposal_rate=3.2, haz_rate=18.0,
        water_source="Groundwater / Borewell",
        leakage_tl=[("2024-09", "Monthly"), ("2025-05", "Frequent")],
        wastewater_tl=[("2024-09", "None")],
        primary_fuel="Electricity Grid", air_control="Basic Scrubber",
        green_tl={
            "ledLighting": [("2024-09", False)], "solarPanels": [("2024-09", False)],
            "rainwaterHarvesting": [("2024-09", False)], "waterRecycling": [("2024-09", False)],
            "wasteSegregation": [("2024-09", True), ("2025-09", False)],
            "energyEfficientMachinery": [("2024-09", False)],
            "evAdoption": [("2024-09", False)], "sustainableMaterials": [("2024-09", False)],
        },
        noise=0.032, seed=505,
        before_after=False, interventions=[],
    ),
]

# Per-state operating-day overrides (festival practice differs by state)
DAY_OVERRIDES = {
    "B001": {"2025-01": 24, "2026-01": 24, "2024-11": 24},                 # TN: Pongal
    "B002": {"2024-11": 21, "2025-10": 23, "2025-11": 25},                 # GJ: Diwali week
    "B003": {"2024-11": 21, "2025-10": 23, "2025-11": 25},
    "B004": {"2024-11": 21, "2025-10": 23, "2025-11": 25},
    "B005": {"2024-11": 21, "2025-10": 23, "2025-11": 25, "2025-03": 25},  # RJ: Diwali + Holi
}


def operating_days(bid, month):
    return DAY_OVERRIDES.get(bid, {}).get(month, BASE_DAYS[month])


# ---------------------------------------------------------------------------
# MONTHLY MODEL
# ---------------------------------------------------------------------------
def build_monthly(b):
    rng = random.Random(b["seed"])
    days = [operating_days(b["bid"], m) for m in MONTHS]
    day_ratio = [d / float(b["days"]) for d in days]

    def scaled(vals):
        return [v * dr for v, dr in zip(vals, day_ratio)]

    electricity = scaled(series(b["kwh_daily"] * b["days"], b["kwh_trend"],
                               b["kwh_seasonal"], b["kwh_breaks"], b["noise"], rng))
    # Production follows the same seasonality as the load but NOT the energy
    # breakpoints (a lighting retrofit must not reduce output).
    production = scaled(series(b["prod_base"], b["kwh_trend"] * 0.6,
                               b["kwh_seasonal"], b.get("prod_breaks"), b["noise"] * 0.8, rng))
    water = scaled(series(b["water_daily"] * b["days"], b["water_trend"],
                          b["water_seasonal"], b.get("water_breaks"), b["noise"], rng))
    waste_streams = {}
    for k, base in b["waste"].items():
        waste_streams[k] = scaled(series(base, b["waste_trend"], SEASON_FLAT, None,
                                         b["noise"] * 1.3, rng))
    genset_diesel = scaled(series(b["genset_diesel"], 0.0, SEASON_FLAT, None, 0.05, rng))
    process_diesel = scaled(series(b["process_diesel"], 0.002, SEASON_FLAT, None, 0.05, rng)) \
        if b["process_diesel"] else [0.0] * 24
    lpg = scaled(series(b["lpg"], 0.001, SEASON_FLAT, None, 0.04, rng)) if b["lpg"] else [0.0] * 24
    fleet_diesel = scaled(series(b["fleet_diesel"], 0.008, b["kwh_seasonal"], None, 0.05, rng))
    fleet_petrol = scaled(series(b["fleet_petrol"], 0.004, SEASON_FLAT, None, 0.05, rng)) \
        if b["fleet_petrol"] else [0.0] * 24
    fleet_cng = scaled(series(b["fleet_cng"], 0.004, SEASON_FLAT, None, 0.05, rng)) \
        if b["fleet_cng"] else [0.0] * 24
    ev_kwh = scaled(series(b["ev_kwh"], 0.006, SEASON_FLAT, None, 0.06, rng)) \
        if b["ev_kwh"] else [0.0] * 24

    rows = []
    for i, m in enumerate(MONTHS):
        solar_kw = flag_timeline(b["solar_timeline"], m)
        # on-site solar generation: kWp x peak-sun-hours x days x performance ratio
        psh = 5.0 if b["state"] in ("Gujarat", "Rajasthan") else 4.8
        renewable = solar_kw * psh * days[i] * 0.78
        renewable = min(renewable, electricity[i] * 0.55)
        rec = dict(
            month=m, days=days[i], employees=b["employees"],
            electricity=electricity[i], renewable=renewable, grid=electricity[i] - renewable,
            solar_kw=solar_kw, production=production[i], water=water[i],
            genset_diesel=genset_diesel[i], process_diesel=process_diesel[i], lpg=lpg[i],
            fleet_diesel=fleet_diesel[i], fleet_petrol=fleet_petrol[i],
            fleet_cng=fleet_cng[i], ev_kwh=ev_kwh[i],
            vehicles=flag_timeline(b["vehicles_tl"], m),
            dry_recovery=flag_timeline(b["dry_recovery_tl"], m),
            organic_recovery=flag_timeline(b["organic_recovery_tl"], m),
            leakage=flag_timeline(b["leakage_tl"], m),
            wastewater=flag_timeline(b["wastewater_tl"], m),
            treat_fraction=flag_timeline(b["treat_fraction_tl"], m),
        )
        # ---- waste partition (single source of truth) ----
        w = {k: v[i] for k, v in waste_streams.items()}
        total_waste = sum(w.values())
        recyclable = (w["plastic"] * 0.95 + w["paper"] * 0.95 + w["metal"] * 0.98
                      + w["material"] * 0.90)
        recyclable = min(recyclable, total_waste - w["organic"] - w["hazardous"])
        recycled = recyclable * rec["dry_recovery"]
        organic_recovered = w["organic"] * rec["organic_recovery"]
        general_bucket = total_waste - recyclable - w["organic"] - w["hazardous"]
        landfill = total_waste - recycled - organic_recovered - w["hazardous"]
        rec.update(dict(
            w_organic=w["organic"], w_plastic=w["plastic"], w_paper=w["paper"],
            w_metal=w["metal"], w_material=w["material"], w_industrial=w["industrial"],
            w_general=w["general"], w_hazardous=w["hazardous"],
            total_waste=total_waste, recyclable=recyclable, recycled=recycled,
            organic_recovered=organic_recovered, general_bucket=general_bucket,
            landfill=landfill,
            recycling_rate=(recycled + organic_recovered) / total_waste * 100.0,
        ))
        # ---- water ----
        process_water = rec["water"] * b["process_share"]
        cleaning_water = rec["water"] * b["cleaning_share"]
        reused = rec["water"] * b["reuse_fraction"]
        ww_generated = rec["water"] * b["ww_fraction"]
        ww_treated = ww_generated * rec["treat_fraction"]
        rec.update(dict(
            process_water=process_water, cleaning_water=cleaning_water,
            other_water=rec["water"] - process_water - cleaning_water,
            reused=reused, ww_generated=ww_generated, ww_treated=ww_treated,
            reuse_rate=reused / (rec["water"] + reused) * 100.0 if (rec["water"] + reused) else 0.0,
        ))
        # ---- materials ----
        raw = sum(b["mats"].values()) * day_ratio[i] * (1 + b["kwh_trend"] * 0.6) ** i
        recycled_input = b["recycled_input"] * day_ratio[i]
        output_kg = rec["production"] * b["output_kg_per_unit"]
        rec.update(dict(
            raw_material=raw, recycled_material=recycled_input,
            virgin_material=raw - recycled_input, output_kg=output_kg,
            rejected=b["rejected"] * day_ratio[i], packaging=b["packaging"] * day_ratio[i],
            recycled_input_pct=recycled_input / raw * 100.0,
            material_efficiency=output_kg / raw * 100.0,
        ))
        # ---- mobility ----
        own_fleet_km = (rec["fleet_diesel"] * b["fleet_mileage"]
                        + rec["fleet_petrol"] * (b["fleet_mileage"] + 3.0)
                        + rec["fleet_cng"] * (b["fleet_mileage"] + 2.0))
        commute_km = b["employees"] * days[i] * b["commute_km_per_emp_day"]
        rec.update(dict(
            own_fleet_km=own_fleet_km, logistics_km=b["logistics_km"] * day_ratio[i],
            commute_km=commute_km, travel_km=b["travel_km"] * day_ratio[i],
            s3_freight_km=b["s3_freight_km"] * day_ratio[i],
            third_party_freight_km=max(0.0, b["logistics_km"] * day_ratio[i] - own_fleet_km),
        ))
        # ---- emissions (app factors) ----
        all_diesel = rec["genset_diesel"] + rec["process_diesel"] + rec["fleet_diesel"]
        e_diesel = all_diesel * EF_DIESEL / 1000.0
        e_petrol = rec["fleet_petrol"] * EF_PETROL / 1000.0
        e_gas = (rec["lpg"] + rec["fleet_cng"]) * EF_GAS / 1000.0
        e_elec = (rec["grid"] + rec["ev_kwh"]) * EF_ELEC / 1000.0
        scope1 = e_diesel + e_petrol + e_gas
        scope2 = e_elec
        s3_materials = sum(qty * day_ratio[i] * S3_MATERIAL_FACTOR[mat]
                           for mat, qty in b["mats"].items()) / 1000.0
        s3_freight = rec["third_party_freight_km"] * S3_FREIGHT_KG_PER_KM / 1000.0
        s3_commute = rec["commute_km"] * b["commute_factor"] / 1000.0
        s3_travel = rec["travel_km"] * S3_TRAVEL_KG_PER_KM / 1000.0
        s3_waste = rec["landfill"] * S3_LANDFILL_KG_PER_KG / 1000.0
        scope3 = s3_materials + s3_freight + s3_commute + s3_travel + s3_waste
        mobility_e = ((rec["fleet_diesel"] * EF_DIESEL + rec["fleet_petrol"] * EF_PETROL
                       + rec["fleet_cng"] * EF_GAS + rec["ev_kwh"] * EF_ELEC) / 1000.0
                      + s3_commute + s3_travel + s3_freight)
        rec.update(dict(
            all_diesel=all_diesel, e_diesel=e_diesel, e_petrol=e_petrol, e_gas=e_gas,
            e_elec=e_elec, scope1=scope1, scope2=scope2, scope3=scope3,
            s3_materials=s3_materials, s3_freight=s3_freight, s3_commute=s3_commute,
            s3_travel=s3_travel, s3_waste=s3_waste,
            total_emissions=scope1 + scope2 + scope3, mobility_emissions=mobility_e,
        ))
        # Round the published base quantities first, then derive every ratio from the
        # rounded values so the published columns reconcile exactly.
        rec["electricity"] = r(rec["electricity"])
        rec["renewable"] = min(r(rec["renewable"]), rec["electricity"])
        rec["grid"] = rec["electricity"] - rec["renewable"]
        rec["production"] = r(rec["production"])
        rec["water"] = r(rec["water"])
        rec["energy_intensity"] = rec["electricity"] / rec["production"]
        rec["water_intensity"] = rec["water"] / rec["production"]
        rec["emissions_intensity"] = rec["total_emissions"] * 1000.0 / rec["production"]
        rec["renewable_share"] = (rec["renewable"] / rec["electricity"] * 100.0
                                  if rec["electricity"] else 0.0)
        rec["avg_load_kw"] = rec["electricity"] / (days[i] * b["hours"])
        rec["peak_demand"] = rec["avg_load_kw"] * b["peak_load_factor"]
        rec["disposal_cost"] = rec["landfill"] * b["disposal_rate"] + rec["w_hazardous"] * b["haz_rate"]
        rec["waste_reduction_pct"] = 0.0  # filled after loop (vs 12 months earlier)
        rows.append(rec)

    for i, rec in enumerate(rows):
        if i >= 12 and rows[i - 12]["total_waste"] > 0:
            rec["waste_reduction_pct"] = (rows[i - 12]["total_waste"] - rec["total_waste"]) \
                / rows[i - 12]["total_waste"] * 100.0
    return rows


def assessment_payload(b, rec):
    """Exact ClimaCred /api/assessment body for one business-month."""
    green = {k: bool(flag_timeline(tl, rec["month"])) for k, tl in b["green_tl"].items()}
    return {
        "energy": {
            "monthlyElectricityKwh": r(rec["electricity"]),
            "monthlyElectricityBillInr": r(rec["electricity"] * b["tariff"]),
            "dieselGeneratorHoursPerMonth": b["genset_hours"],
            "generatorFuelLitresPerMonth": r(rec["genset_diesel"]),
            "existingSolarCapacityKw": rec["solar_kw"],
            "energyEfficientEquipmentPercent": b["eff_equip"],
        },
        "water": {
            "monthlyWaterLitres": r(rec["water"]),
            "waterSource": b["water_source"],
            "waterRecyclingAvailable": green["waterRecycling"],
            "rainwaterHarvesting": green["rainwaterHarvesting"],
            "leakageFrequency": rec["leakage"],
            "wastewaterTreatment": rec["wastewater"],
        },
        "waste": {
            "organicWasteKgPerMonth": r(rec["w_organic"]),
            "plasticWasteKgPerMonth": r(rec["w_plastic"]),
            "paperWasteKgPerMonth": r(rec["w_paper"]),
            "industrialWasteKgPerMonth": r(rec["w_industrial"] + rec["w_general"] + rec["w_hazardous"]),
            "textileMaterialWasteKgPerMonth": r(rec["w_material"]),
            "currentRecyclingPercent": r(rec["recycling_rate"], 1),
            "wasteSegregationPracticed": green["wasteSegregation"],
        },
        "emissions": {
            "primaryFuel": b["primary_fuel"],
            "monthlyDieselLitres": r(rec["all_diesel"]),
            "monthlyPetrolLitres": r(rec["fleet_petrol"]),
            "monthlyNaturalGasKg": r(rec["lpg"] + rec["fleet_cng"]),
            "mainEmissionSources": ["Grid Electricity", "Diesel Genset", "Fleet Transport",
                                     "Process Heat", "Purchased Materials"],
            "airPollutionControlSystem": b["air_control"],
        },
        "mobility": {
            "deliveryVehiclesCount": int(rec["vehicles"]),
            "vehicleFuelType": b["vehicle_fuel"],
            "monthlyFleetFuelLitres": r(rec["fleet_diesel"] + rec["fleet_petrol"]),
            "employeeCommuteMode": b["commute_mode"],
            "evAdoptedPercent": r(100.0 / rec["vehicles"]) if b["ev_kwh"] else 0.0,
        },
        "greenPractices": green,
    }


def profile_payload(b):
    """Exact ClimaCred /api/profile body."""
    return {
        "business_name": b["name"], "name": b["name"],
        "industry": b["industry"], "business_type": b["sub_industry"],
        "businessType": b["sub_industry"],
        "location": f"{b['city']}, {b['state']}, India",
        "employees": b["employees"],
        "working_days": b["days"], "workingDaysPerMonth": b["days"],
        "production_volume": f"{b['prod_base']:,} {b['prod_unit']} / month",
        "productionVolume": f"{b['prod_base']:,} {b['prod_unit']} / month",
        "operating_hours": b["hours"], "operatingHoursPerDay": b["hours"],
        "business_size": b["business_size"], "businessSize": b["business_size"],
        "facility_area_sqft": b["facility_area_sqft"],
        "facilityAreaSqFt": b["facility_area_sqft"],
        "contact_email": b["email"], "contactEmail": b["email"],
        "phone": "+91 00000 00000 (placeholder)",
    }


# ---------------------------------------------------------------------------
# Build everything
# ---------------------------------------------------------------------------
DATA = {}
for b in BUSINESSES:
    rows = build_monthly(b)
    fps = []
    for rec in rows:
        fps.append(generate_fingerprint(profile_payload(b), assessment_payload(b, rec)))
    DATA[b["bid"]] = dict(b=b, rows=rows, fps=fps,
                          profile=profile_payload(b), assessment=assessment_payload(b, rows[-1]))
    DATA[b["bid"]]["fingerprint"] = fps[-1]

# ---------------------------------------------------------------------------
# SOLUTION CATALOG: the 12 in-app solutions (exact ids/values) + 6 extended
# ---------------------------------------------------------------------------
EXTENDED_SOLUTIONS = [
    dict(id="sol-etp-optimization", name="Effluent Treatment Plant Optimization",
         category="Water", capex=(850000, 1150000), savings=310000, kwh=12000,
         water=420000, waste=0, months=6, complexity="Medium", life=12,
         desc="Aeration tuning, dosing automation and treated-water return line so ETP outflow can be reused in cooling and floor washing.",
         industries="Textile; Food Processing; Chemicals; Packaging"),
    dict(id="sol-waste-segregation", name="Source Waste Segregation & Colour-Coded Bins",
         category="Waste", capex=(90000, 160000), savings=185000, kwh=0,
         water=0, waste=24000, months=2, complexity="Low", life=8,
         desc="Four-stream segregation at source with weighbridge logging, operator training and vendor-wise scrap routing.",
         industries="All"),
    dict(id="sol-composting-biogas", name="Organic Waste Composting / Small Biogas Unit",
         category="Waste", capex=(650000, 950000), savings=275000, kwh=0,
         water=0, waste=48000, months=5, complexity="Medium", life=10,
         desc="Windrow composting or a small plug-flow digester for food/organic residue, replacing landfill disposal and part of the LPG load.",
         industries="Food Processing; Textile (canteen); Hospitality"),
    dict(id="sol-sustainable-packaging", name="Sustainable / Light-Weight Packaging Design",
         category="Materials", capex=(420000, 680000), savings=230000, kwh=0,
         water=0, waste=18000, months=4, complexity="Medium", life=6,
         desc="Down-gauging, recycled-board substitution and reusable transport packaging to cut virgin polymer and fibre input.",
         industries="Packaging; Food Processing; Textile; Furniture"),
    dict(id="sol-hvac-optimization", name="HVAC & Cold-Chain Optimization",
         category="Energy", capex=(520000, 780000), savings=265000, kwh=32000,
         water=0, waste=0, months=5, complexity="Medium", life=12,
         desc="Set-point reset, door curtains/air locks, condenser cleaning schedule and VFD-controlled cold-room compressors.",
         industries="Food Processing; Textile; Precision Engineering"),
    dict(id="sol-compressed-air", name="Compressed Air System Optimization",
         category="Energy", capex=(340000, 520000), savings=245000, kwh=29800,
         water=0, waste=0, months=3, complexity="Low", life=10,
         desc="Ultrasonic leak survey, pressure-band reduction, receiver sizing and heat-of-compression recovery.",
         industries="Precision Engineering; Textile; Packaging; Furniture"),
]


def solution_rows():
    rows = []
    for s in SOLUTION_CATALOG:
        mid = (s["investmentMinInr"] + s["investmentMaxInr"]) / 2.0
        rows.append(dict(
            solution_id=s["id"], solution_name=s["title"], category=s["category"],
            description=s["description"],
            applicable_industries=", ".join(s.get("applicable_industries") or ["All (reference sizing)"]),
            capex_min=s["investmentMinInr"], capex_max=s["investmentMaxInr"],
            capex_mid=mid, savings=s["potentialAnnualSavingsInr"],
            co2=s["co2ReductionTonnesPerYear"], water_l=0, waste_kg=0,
            payback=mid / s["potentialAnnualSavingsInr"] * 12.0,
            impl_months={"Low": 3, "Medium": 6, "High": 10}[s["implementationDifficulty"]],
            complexity=s["implementationDifficulty"], life=15,
            in_app="Yes", assumptions="; ".join(s.get("assumptions", [])),
        ))
    # water / waste saving columns for the in-app water & waste solutions
    water_map = {"sol-water-ro": 3740000, "sol-rainwater": 980000, "sol-leak-sensors": 420000}
    waste_map = {"sol-waste-recovery": 28500, "sol-material-reuse": 21000}
    for row in rows:
        row["water_l"] = water_map.get(row["solution_id"], 0)
        row["waste_kg"] = waste_map.get(row["solution_id"], 0)
        if row["solution_id"] in ("sol-solar", "sol-machinery-vfd", "sol-led-iot",
                                  "sol-smart-metering", "sol-heat-recovery"):
            row["kwh_saving"] = row["co2"] * 1000.0 / EF_ELEC
        else:
            row["kwh_saving"] = 0
    for s in EXTENDED_SOLUTIONS:
        mid = sum(s["capex"]) / 2.0
        co2 = s["kwh"] * EF_ELEC / 1000.0
        if s["id"] == "sol-waste-segregation":
            co2 = s["waste"] * S3_LANDFILL_KG_PER_KG / 1000.0
        if s["id"] == "sol-composting-biogas":
            co2 = s["waste"] * 0.90 / 1000.0 + 60 * 12 * EF_GAS / 1000.0
        if s["id"] == "sol-sustainable-packaging":
            co2 = s["waste"] * (S3_MATERIAL_FACTOR["polymer"] - S3_MATERIAL_FACTOR["kraft_paper"]) / 1000.0
        if s["id"] == "sol-etp-optimization":
            co2 = s["kwh"] * EF_ELEC / 1000.0
        rows.append(dict(
            solution_id=s["id"], solution_name=s["name"], category=s["category"],
            description=s["desc"], applicable_industries=s["industries"],
            capex_min=s["capex"][0], capex_max=s["capex"][1], capex_mid=mid,
            savings=s["savings"], co2=co2, water_l=s["water"], waste_kg=s["waste"],
            payback=mid / s["savings"] * 12.0, impl_months=s["months"],
            complexity=s["complexity"], life=s["life"], in_app="No",
            assumptions="Illustrative estimate for testing only; sizing must be re-validated on site.",
            kwh_saving=s["kwh"],
        ))
    return rows


SOLUTIONS = {s["solution_id"]: s for s in solution_rows()}

# ---------------------------------------------------------------------------
# Scaling of catalog economics to each business (documented in README)
# ---------------------------------------------------------------------------
REF = dict(energy=265000.0, water=4600000.0, waste=17970.0, mobility=2850.0,
           materials=76500.0, area=42000.0)
SCALE_BOUNDS = (0.35, 1.6)


def scale_for(b, rec, category):
    cur = DATA[b["bid"]]["rows"][-1]
    if category == "Energy":
        raw = cur["electricity"] / REF["energy"]
    elif category == "Water":
        raw = cur["water"] / REF["water"]
    elif category == "Waste":
        raw = cur["total_waste"] / REF["waste"]
    elif category == "Mobility":
        raw = (cur["fleet_diesel"] + cur["fleet_petrol"]) / REF["mobility"]
    elif category == "Materials":
        raw = cur["raw_material"] / REF["materials"]
    else:
        raw = b["facility_area_sqft"] / REF["area"]
    return max(SCALE_BOUNDS[0], min(SCALE_BOUNDS[1], raw))


def scaled_solution(b, rec, sid):
    sol = SOLUTIONS[sid]
    f = scale_for(b, rec, sol["category"])
    capex = int(round(sol["capex_mid"] * f / 1000.0) * 1000)
    savings = int(round(sol["savings"] * f / 1000.0) * 1000)
    co2 = sol["co2"] * f
    kwh = sol["kwh_saving"] * f
    water_l = sol["water_l"] * f
    waste_kg = sol["waste_kg"] * f
    return dict(solution_id=sid, name=sol["solution_name"], category=sol["category"],
                capex=capex, savings=savings, co2=co2, kwh=kwh,
                water_l=water_l, waste_kg=waste_kg,
                payback_months=capex / savings * 12.0 if savings else 0.0,
                impl_months=sol["impl_months"], complexity=sol["complexity"], scale=f)


# ---------------------------------------------------------------------------
# Recommendation rules (derived from each business's own stored numbers)
# ---------------------------------------------------------------------------
def build_recommendations(b):
    cur = DATA[b["bid"]]["rows"][-1]
    green = DATA[b["bid"]]["assessment"]["greenPractices"]
    cands = []  # (solution_id, reason)

    if cur["electricity"] > 30000 and cur["renewable_share"] < 30:
        existing = (f"Existing capacity is {cur['solar_kw']:.0f} kWp."
                    if cur["solar_kw"] else "No on-site array is recorded.")
        cands.append(("sol-solar",
                      f"On-site renewables supply only {cur['renewable_share']:.1f}% of "
                      f"{cur['electricity']:,.0f} kWh/month of grid electricity. {existing}"))
    if b["eff_equip"] < 55 and cur["electricity"] > 30000:
        cands.append(("sol-machinery-vfd",
                      f"Only {b['eff_equip']:.0f}% of equipment is energy-efficient while the plant "
                      f"draws {cur['electricity']:,.0f} kWh/month "
                      f"({cur['energy_intensity']:.2f} kWh per {b['prod_unit']})."))
    if not green["ledLighting"]:
        cands.append(("sol-led-iot",
                      "LED lighting is not recorded as adopted; high-bay lighting runs across "
                      f"{b['hours']} operating hours/day."))
    if b["eff_equip"] < 80:
        cands.append(("sol-smart-metering",
                      "No departmental sub-metering is recorded, so the "
                      f"{cur['electricity']:,.0f} kWh/month load cannot be attributed to lines."))
    if cur["water"] > 200000 and cur["reuse_rate"] < 10:
        cands.append(("sol-water-ro",
                      f"Freshwater intake is {cur['water']:,.0f} L/month "
                      f"({cur['water_intensity']:.1f} L per {b['prod_unit']}) with only "
                      f"{cur['reuse_rate']:.1f}% reuse."))
    if cur["leakage"] in ("Monthly", "Frequent"):
        cands.append(("sol-leak-sensors",
                      f"Leakage frequency is reported as '{cur['leakage']}' on "
                      f"{cur['water']:,.0f} L/month of intake."))
    if not green["rainwaterHarvesting"] and b["facility_area_sqft"] >= 15000:
        cands.append(("sol-rainwater",
                      f"No rainwater harvesting on a {b['facility_area_sqft']:,} sq ft roof in "
                      f"{b['city']}, {b['state']}."))
    if cur["treat_fraction"] < 0.85 and cur["ww_generated"] > 20000:
        cands.append(("sol-etp-optimization",
                      f"Only {cur['treat_fraction']*100:.0f}% of the "
                      f"{cur['ww_generated']:,.0f} L/month of wastewater is treated "
                      f"(treatment level: {cur['wastewater']})."))
    if cur["w_organic"] > 3000:
        cands.append(("sol-composting-biogas",
                      f"Organic residue is {cur['w_organic']:,.0f} kg/month and only "
                      f"{cur['organic_recovery']*100:.0f}% of it is recovered."))
    if cur["recycling_rate"] < 50:
        cands.append(("sol-waste-segregation",
                      f"Recycling rate is {cur['recycling_rate']:.1f}% even though "
                      f"{cur['recyclable']:,.0f} kg/month of the "
                      f"{cur['total_waste']:,.0f} kg/month stream is recyclable."))
    if cur["w_material"] > 2000 or cur["w_metal"] > 2000:
        cands.append(("sol-waste-recovery",
                      f"Material scrap ({cur['w_material']:,.0f} kg) and metal scrap "
                      f"({cur['w_metal']:,.0f} kg) leave the site each month."))
    if (cur["fleet_diesel"] + cur["fleet_petrol"]) > 1200 or cur["logistics_km"] > 3000:
        cands.append(("sol-route-opt",
                      f"Own fleet burns {cur['fleet_diesel'] + cur['fleet_petrol']:,.0f} L/month "
                      f"across {int(cur['vehicles'])} vehicles with "
                      f"{cur['logistics_km']:,.0f} km of freight movement."))
    if (cur["fleet_diesel"] + cur["fleet_petrol"]) > 1800 and cur["ev_kwh"] == 0:
        cands.append(("sol-ev-fleet",
                      f"{int(cur['vehicles'])} vehicles are 100% fossil-fuelled "
                      f"({cur['fleet_diesel']:,.0f} L diesel/month)."))
    if cur["recycled_input_pct"] < 20 and b["industry"] != "Food Processing":
        cands.append(("sol-material-reuse",
                      f"Recycled material input is only {cur['recycled_input_pct']:.1f}% of the "
                      f"{cur['raw_material']:,.0f} kg/month material input."))
    if cur["packaging"] > 2000 or b["industry"] == "Packaging Manufacturing":
        cands.append(("sol-sustainable-packaging",
                      f"Packaging material use is {cur['packaging']:,.0f} kg/month with "
                      f"{cur['recycled_input_pct']:.1f}% recycled input."))
    if b["industry"] == "Precision Engineering / Machining":
        cands.append(("sol-compressed-air",
                      f"Machining load of {cur['electricity']:,.0f} kWh/month with "
                      f"{b['eff_equip']:.0f}% efficient equipment points to compressed-air losses."))
    if b["industry"] == "Food Processing":
        cands.append(("sol-hvac-optimization",
                      f"Cold-chain driven load of {cur['electricity']:,.0f} kWh/month "
                      f"({cur['energy_intensity']:.2f} kWh per {b['prod_unit']})."))
    if b["industry"] == "Textile Manufacturing" and cur["genset_diesel"] > 500:
        cands.append(("sol-heat-recovery",
                      f"Thermal demand with {cur['genset_diesel']:,.0f} L/month of genset diesel and "
                      f"{cur['lpg']:,.0f} kg/month of LPG."))

    # Guarantee a minimum of 3 candidates so every business is testable.
    if len({c[0] for c in cands}) < 3:
        if cur["lpg"] + cur["genset_diesel"] > 100:
            cands.append(("sol-heat-recovery",
                          f"Thermal load with {cur['genset_diesel']:,.0f} L/month of genset diesel "
                          f"and {cur['lpg']:,.0f} kg/month of LPG."))
        cands.append(("sol-waste-segregation",
                      f"Recycling rate {cur['recycling_rate']:.1f}% with "
                      f"{cur['recyclable']:,.0f} kg/month recyclable in the "
                      f"{cur['total_waste']:,.0f} kg/month stream."))

    seen, out = set(), []
    for sid, reason in cands:
        if sid in seen:
            continue
        seen.add(sid)
        rec = scaled_solution(b, None, sid)
        rec["reason"] = reason
        out.append(rec)

    # Priority = severity of the dimension the solution addresses + financial
    # opportunity + environmental opportunity (documented weights).
    dims = {d["dimension"]: d["score"] for d in DATA[b["bid"]]["fingerprint"]["dimensions"]}
    dim_for = {"Energy": "Energy", "Water": "Water", "Waste": "Waste", "Mobility": "Mobility",
               "Materials": "Operations", "Operations": "Operations"}
    pbs = [o["payback_months"] for o in out]
    co2s = [o["co2"] for o in out]
    pb_min, pb_max = min(pbs), max(pbs)
    co2_max = max(co2s) or 1.0
    for o in out:
        severity = 1.0 - dims.get(dim_for.get(o["category"], "Operations"), 100) / 100.0
        fin = ((pb_max - o["payback_months"]) / (pb_max - pb_min)) if pb_max > pb_min else 1.0
        env = o["co2"] / co2_max
        o["severity"] = severity
        o["score"] = 0.50 * severity + 0.30 * fin + 0.20 * env
    out.sort(key=lambda x: (-x["score"], x["payback_months"]))
    out = out[:6]
    for i, rec in enumerate(out, start=1):
        rec["priority"] = i
        rec["addresses_dimension"] = dim_for.get(rec["category"], "Operations")
        rec["dimension_score"] = dims.get(rec["addresses_dimension"], 0)
        rec["timeline"] = f"Month {1 + (i - 1) * 2} - {rec['impl_months'] + (i - 1) * 2}"
        rec["improvement"] = (
            f"{rec['savings']/1000.0:.1f}k INR/yr indicative saving; "
            f"{rec['co2']:.1f} tCO2e/yr indicative reduction")
    return out


RECOMMENDATIONS = {b["bid"]: build_recommendations(b) for b in BUSINESSES}

# ---------------------------------------------------------------------------
# Scenarios
# ---------------------------------------------------------------------------
SCENARIO_NAMES = ["Current Baseline", "Solar Scenario", "Energy Efficiency Scenario",
                  "Water Efficiency Scenario", "Waste Circularity Scenario",
                  "Green Mobility Scenario", "Combined Transformation Scenario"]


def scenario_solutions(b, name):
    """Applicable solutions for one scenario, filtered to this business's own data.

    The Combined Transformation Scenario is the union of the other six, so by
    construction it is always at least as large as any single scenario.
    """
    cur = DATA[b["bid"]]["rows"][-1]
    if name == "Current Baseline":
        return []
    if name == "Solar Scenario":
        return ["sol-solar"]
    if name == "Energy Efficiency Scenario":
        out = ["sol-machinery-vfd", "sol-led-iot", "sol-smart-metering"]
        if b["industry"] in ("Precision Engineering / Machining", "Packaging Manufacturing"):
            out.append("sol-compressed-air")
        if b["industry"] == "Food Processing" or b["eff_equip"] < 40:
            out.append("sol-hvac-optimization")
        return out
    if name == "Water Efficiency Scenario":
        out = []
        if cur["water"] > 20000:
            out.append("sol-leak-sensors")
        if b["facility_area_sqft"] >= 15000:
            out.append("sol-rainwater")
        if cur["reuse_rate"] < 20 and cur["water"] > 100000:
            out.append("sol-water-ro")
        if cur["treat_fraction"] < 0.85 and cur["ww_generated"] > 20000:
            out.append("sol-etp-optimization")
        return out
    if name == "Waste Circularity Scenario":
        out = []
        if cur["recycling_rate"] < 60:
            out.append("sol-waste-segregation")
        if cur["w_material"] > 2000 or cur["w_metal"] > 2000:
            out.append("sol-waste-recovery")
        if cur["w_organic"] > 3000:
            out.append("sol-composting-biogas")
        if cur["packaging"] > 2000:
            out.append("sol-sustainable-packaging")
        return out
    if name == "Green Mobility Scenario":
        fuel = cur["fleet_diesel"] + cur["fleet_petrol"]
        out = []
        if fuel > 300:
            out.append("sol-route-opt")
        if fuel > 1200:
            out.append("sol-ev-fleet")
        return out
    # Combined: union of every other scenario, in a stable order
    union = []
    for other in SCENARIO_NAMES[1:-1]:
        for sid in scenario_solutions(b, other):
            if sid not in union:
                union.append(sid)
    return union


TRANSFORMATION_SET = {b["bid"]: scenario_solutions(b, "Combined Transformation Scenario")
                      for b in BUSINESSES}
OM_RATE = {"sol-solar": 0.03, "sol-water-ro": 0.06, "sol-etp-optimization": 0.05,
           "sol-smart-metering": 0.02, "sol-ev-fleet": 0.04, "sol-composting-biogas": 0.04,
           "sol-hvac-optimization": 0.03}


def scenario_rows():
    rows = []
    for b in BUSINESSES:
        rec_ids = [x["solution_id"] for x in RECOMMENDATIONS[b["bid"]]]
        cur = DATA[b["bid"]]["rows"][-1]
        for n, name in enumerate(SCENARIO_NAMES, start=1):
            sids = scenario_solutions(b, name)
            picked = [scaled_solution(b, None, s) for s in sids if s in SOLUTIONS]
            capex = sum(p["capex"] for p in picked)
            savings = sum(p["savings"] for p in picked)
            om = sum(p["capex"] * OM_RATE.get(p["solution_id"], 0.0) for p in picked)
            rows.append(dict(
                scenario_id=f"SCN-{b['bid']}-{n:02d}", business_id=b["bid"], scenario_name=name,
                solutions=", ".join(p["solution_id"] for p in picked) or "None (business as usual)",
                capex=r(capex), om=r(om),
                energy_kwh=r(sum(p["kwh"] for p in picked)),
                water_l=r(sum(p["water_l"] for p in picked)),
                waste_kg=r(sum(p["waste_kg"] for p in picked)),
                co2=r(sum(p["co2"] for p in picked), 2),
                savings=r(savings),
                payback_years=r(capex / savings, 2) if savings else 0.0,
                five_year=r(5 * (savings - om) - capex),
                impl_months=max([p["impl_months"] for p in picked], default=0),
                baseline_kwh=r(cur["electricity"] * 12),
                baseline_water=r(cur["water"] * 12),
                baseline_co2=r(cur["total_emissions"] * 12, 2),
            ))
    return rows


SCENARIOS = scenario_rows()

# ---------------------------------------------------------------------------
# Transformation plan
# ---------------------------------------------------------------------------
PLAN_PHASES = {
    "sol-waste-segregation": (1, "Phase 1 - Quick Wins"),
    "sol-route-opt": (1, "Phase 1 - Quick Wins"),
    "sol-leak-sensors": (1, "Phase 1 - Quick Wins"),
    "sol-led-iot": (1, "Phase 1 - Quick Wins"),
    "sol-machinery-vfd": (2, "Phase 2 - Efficiency Improvements"),
    "sol-smart-metering": (2, "Phase 2 - Efficiency Improvements"),
    "sol-compressed-air": (2, "Phase 2 - Efficiency Improvements"),
    "sol-hvac-optimization": (2, "Phase 2 - Efficiency Improvements"),
    "sol-rainwater": (2, "Phase 2 - Efficiency Improvements"),
    "sol-etp-optimization": (2, "Phase 2 - Efficiency Improvements"),
    "sol-waste-recovery": (3, "Phase 3 - Circularity"),
    "sol-composting-biogas": (3, "Phase 3 - Circularity"),
    "sol-material-reuse": (3, "Phase 3 - Circularity"),
    "sol-sustainable-packaging": (3, "Phase 3 - Circularity"),
    "sol-solar": (4, "Phase 4 - Renewable / Mobility Transition"),
    "sol-ev-fleet": (4, "Phase 4 - Renewable / Mobility Transition"),
    "sol-heat-recovery": (4, "Phase 4 - Renewable / Mobility Transition"),
}
OWNER = {"Energy": "Energy Manager", "Water": "EHS Officer", "Waste": "EHS Officer",
         "Mobility": "Logistics Manager", "Materials": "Procurement Lead",
         "Operations": "Plant Head"}
DEPENDENCY = {
    "sol-solar": "Roof structural survey + DISCOM net-metering application",
    "sol-water-ro": "ETP stabilisation before return line commissioning",
    "sol-machinery-vfd": "Sub-metering baseline from sol-smart-metering preferred",
    "sol-etp-optimization": "Wastewater flow survey",
    "sol-ev-fleet": "Charging infrastructure and depot power sanction",
    "sol-composting-biogas": "Segregation at source must be running",
    "sol-waste-recovery": "Segregation at source must be running",
    "sol-material-reuse": "Buyer specification approval for recycled content",
    "sol-sustainable-packaging": "Customer drop-test / specification approval",
    "sol-heat-recovery": "Boiler / thermic-fluid efficiency audit",
}


def plan_rows():
    rows = []
    for b in BUSINESSES:
        recs = [scaled_solution(b, None, sid) for sid in TRANSFORMATION_SET[b["bid"]]]
        has_ba = b["before_after"]
        items = []
        for rec in recs:
            ph, ph_name = PLAN_PHASES.get(rec["solution_id"], (2, "Phase 2 - Efficiency Improvements"))
            items.append((ph, ph_name, rec))
        items.sort(key=lambda x: (x[0], x[2]["payback_months"]))
        start = 1
        cursor = {1: 1, 2: 4, 3: 8, 4: 12, 5: 15}
        for ph, ph_name, rec in items:
            if has_ba:
                status = "Completed" if ph == 1 else ("In Progress" if ph == 2 else "Planned")
            else:
                status = "In Progress" if (ph == 1 and start == 1) else "Planned"
            rows.append(dict(
                business_id=b["bid"], phase=ph_name, action=rec["name"],
                solution_id=rec["solution_id"], start_month=cursor[ph],
                duration_months=rec["impl_months"], expected_cost=r(rec["capex"]),
                expected_savings=r(rec["savings"]),
                expected_emission_reduction=r(rec["co2"], 2),
                expected_water_saving=r(rec["water_l"]),
                owner_role=OWNER.get(rec["category"], "Plant Head"),
                dependency=DEPENDENCY.get(rec["solution_id"], "Baseline data from Climate Assessment"),
                status=status,
            ))
            cursor[ph] += rec["impl_months"]
        rows.append(dict(
            business_id=b["bid"], phase="Phase 5 - Verification",
            action="Metering calibration, data reconciliation & third-party review of reported outcomes",
            solution_id="N/A", start_month=15, duration_months=3, expected_cost=180000,
            expected_savings=0, expected_emission_reduction=0.0, expected_water_saving=0,
            owner_role="Sustainability Coordinator",
            dependency="At least 6 months of post-implementation meter readings",
            status="Planned" if not has_ba else "In Progress",
        ))
    return rows


PLAN = plan_rows()

# ---------------------------------------------------------------------------
# Before / after impact verification (3 businesses, derived from the model)
# ---------------------------------------------------------------------------
BEFORE_WINDOW = ["2025-03", "2025-04", "2025-05", "2025-06", "2025-07", "2025-08"]
AFTER_WINDOW = MONTHS[-6:]


def window_avg(bid, key, window):
    vals = [rec[key] for rec in DATA[bid]["rows"] if rec["month"] in window]
    return sum(vals) / len(vals)


def impact_rows():
    rows = []
    for b in BUSINESSES:
        if not b["before_after"]:
            continue
        bid = b["bid"]
        metrics = [
            ("Electricity consumption", "electricity", "kWh/month", 0, True),
            ("Water consumption", "water", "litres/month", 0, True),
            ("Waste to landfill", "landfill", "kg/month", 0, True),
            ("Recycling rate", "recycling_rate", "%", 1, False),
            ("Diesel consumption", "all_diesel", "litres/month", 0, True),
            ("Scope 1 emissions", "scope1", "tCO2e/month", 2, True),
            ("Scope 2 emissions", "scope2", "tCO2e/month", 2, True),
            ("Total emissions", "total_emissions", "tCO2e/month", 2, True),
            ("Energy intensity", "energy_intensity", "kWh per production unit", 3, True),
            ("Water intensity", "water_intensity", "litres per production unit", 2, True),
        ]
        for label, key, unit, nd, lower_is_better in metrics:
            base = r(window_avg(bid, key, BEFORE_WINDOW), nd)
            after = r(window_avg(bid, key, AFTER_WINDOW), nd)
            # positive = favourable: a fall in a lower-is-better metric, or a rise
            # in a higher-is-better metric (e.g. recycling rate).
            if base:
                pct = ((base - after) if lower_is_better else (after - base)) / base * 100.0
            else:
                pct = 0.0
            rows.append(dict(
                business_id=bid, metric=label,
                baseline_value=base, after_value=after, unit=unit,
                improvement_percent=r(pct, 1),
                verification_status="Self-reported by business (not third-party verified)",
                measurement_period=f"Baseline {BEFORE_WINDOW[0]} to {BEFORE_WINDOW[-1]} vs "
                                   f"Post {AFTER_WINDOW[0]} to {AFTER_WINDOW[-1]}",
                direction=("Reduction is favourable" if lower_is_better else "Increase is favourable"),
            ))
    return rows


IMPACT = impact_rows()

# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)


def write_xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for c, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(rw[c - 1])) for rw in rows[:200]] or [10])
            ws.column_dimensions[get_column_letter(c)].width = min(max(11, width + 2), 58)
    wb.save(path)


def write_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(headers)
        w.writerows(rows)


# ---- 1. business_profiles.xlsx ---------------------------------------------
profile_headers = ["business_id", "business_name", "industry", "sub_industry", "city", "state",
                   "employee_count", "facility_area_sqft", "operating_days_per_month",
                   "operating_hours_per_day", "annual_turnover", "ownership_type",
                   "facility_type", "year_established", "primary_products", "contact_email",
                   "created_at", "business_size", "test_archetype"]
profile_rows = [[b["bid"], b["name"], b["industry"], b["sub_industry"], b["city"], b["state"],
                 b["employees"], b["facility_area_sqft"], b["days"], b["hours"], b["turnover"],
                 b["ownership"], b["facility_type"], b["established"], b["products"],
                 b["email"], b["created_at"], b["business_size"], b["archetype"]]
                for b in BUSINESSES]
profile_api = [[b["bid"], json.dumps(DATA[b["bid"]]["profile"], ensure_ascii=False)]
               for b in BUSINESSES]

# ---- 2. climate_assessments.xlsx -------------------------------------------
assess_headers = ["assessment_id", "business_id", "assessment_date", "energy_score", "water_score",
                  "waste_score", "emissions_score", "mobility_score", "operations_score",
                  "overall_climate_score", "risk_level", "key_issue_1", "key_issue_2",
                  "key_issue_3", "energy_impact_level", "water_impact_level", "waste_impact_level",
                  "emissions_impact_level", "mobility_impact_level", "operations_impact_level",
                  "data_completeness_percent", "scoring_source"]
assess_rows = []
assess_api = []
for b in BUSINESSES:
    fp = DATA[b["bid"]]["fingerprint"]
    dims = {d["dimension"]: d for d in fp["dimensions"]}
    top = fp["top_improvement_dimensions"]
    issues = [f"{d}: {dims[d]['primary_cause']}" for d in top]
    issues += ["", "", ""]
    assess_rows.append([
        f"CA-{b['bid']}-{CURRENT_MONTH}", b["bid"], ASSESSMENT_DATE,
        dims["Energy"]["score"], dims["Water"]["score"], dims["Waste"]["score"],
        dims["Emissions"]["score"], dims["Mobility"]["score"], dims["Operations"]["score"],
        fp["overallScore"], fp["scoreLabel"], issues[0], issues[1], issues[2],
        dims["Energy"]["impact_level"], dims["Water"]["impact_level"], dims["Waste"]["impact_level"],
        dims["Emissions"]["impact_level"], dims["Mobility"]["impact_level"],
        dims["Operations"]["impact_level"], fp["data_quality"]["completeness_percent"],
        "Computed by backend/app/climate_engine (generate_fingerprint) on the assessment payload",
    ])
    app_em = calculate_emissions_metrics(DATA[b["bid"]]["assessment"])
    assess_api.append([
        b["bid"], json.dumps(DATA[b["bid"]]["assessment"], ensure_ascii=False),
        app_em["emissions_breakdown_tonnes_co2e_per_month"]["total"],
        r(DATA[b["bid"]]["rows"][-1]["total_emissions"], 2),
        r(DATA[b["bid"]]["rows"][-1]["scope1"] + DATA[b["bid"]]["rows"][-1]["scope2"], 2),
        fp["overallScore"],
    ])

# ---- 3. resource_consumption.xlsx (last 12 months) --------------------------
resource_headers = ["business_id", "month", "electricity_kwh", "diesel_litres", "LPG_kg",
                    "water_litres", "wastewater_litres", "total_waste_kg", "recyclable_waste_kg",
                    "hazardous_waste_kg", "organic_waste_kg", "general_waste_kg",
                    "production_units", "operating_days", "employee_count"]
resource_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        resource_rows.append([
            b["bid"], rec["month"], r(rec["electricity"]), r(rec["all_diesel"]), r(rec["lpg"]),
            r(rec["water"]), r(rec["ww_generated"]), r(rec["total_waste"]), r(rec["recyclable"]),
            r(rec["w_hazardous"]), r(rec["w_organic"]), r(rec["general_bucket"]),
            r(rec["production"]), rec["days"], rec["employees"]])

# ---- 4. emissions_data.xlsx -------------------------------------------------
emissions_headers = ["business_id", "month", "scope_1_emissions_tco2e", "scope_2_emissions_tco2e",
                     "scope_3_emissions_tco2e", "total_emissions_tco2e",
                     "diesel_emissions_tco2e", "electricity_emissions_tco2e",
                     "mobility_emissions_tco2e", "emissions_intensity"]
emissions_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        emissions_rows.append([
            b["bid"], rec["month"], r(rec["scope1"], 3), r(rec["scope2"], 3),
            r(rec["scope3"], 3), r(rec["total_emissions"], 3), r(rec["e_diesel"], 3),
            r(rec["e_elec"], 3), r(rec["mobility_emissions"], 3),
            r(rec["emissions_intensity"], 4)])

# ---- 5. mobility_data.xlsx --------------------------------------------------
mobility_headers = ["business_id", "month", "company_vehicle_count", "diesel_litres",
                    "petrol_litres", "CNG_kg", "EV_kwh", "employee_commute_km",
                    "business_travel_km", "logistics_km", "estimated_mobility_emissions_tco2e"]
mobility_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        mobility_rows.append([
            b["bid"], rec["month"], int(rec["vehicles"]), r(rec["fleet_diesel"]),
            r(rec["fleet_petrol"]), r(rec["fleet_cng"]), r(rec["ev_kwh"]), r(rec["commute_km"]),
            r(rec["travel_km"]), r(rec["logistics_km"]), r(rec["mobility_emissions"], 3)])

# ---- 6. waste_data.xlsx -----------------------------------------------------
waste_headers = ["business_id", "month", "total_waste_kg", "recyclable_kg", "recycled_kg",
                 "organic_kg", "organic_recovered_kg", "hazardous_kg", "landfill_kg",
                 "waste_reduction_percent", "recycling_rate_percent", "disposal_cost"]
waste_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        waste_rows.append([
            b["bid"], rec["month"], r(rec["total_waste"]), r(rec["recyclable"]),
            r(rec["recycled"]), r(rec["w_organic"]), r(rec["organic_recovered"]),
            r(rec["w_hazardous"]), r(rec["landfill"]),
            r(rec["waste_reduction_pct"], 1), r(rec["recycling_rate"], 1),
            r(rec["disposal_cost"])])

# ---- 7. water_data.xlsx -----------------------------------------------------
water_headers = ["business_id", "month", "freshwater_intake_litres", "process_water_litres",
                 "cleaning_water_litres", "reused_water_litres", "wastewater_generated_litres",
                 "wastewater_treated_litres", "water_reuse_rate_percent",
                 "water_intensity_litres_per_unit"]
water_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        water_rows.append([
            b["bid"], rec["month"], r(rec["water"]), r(rec["process_water"]),
            r(rec["cleaning_water"]), r(rec["reused"]), r(rec["ww_generated"]),
            r(rec["ww_treated"]), r(rec["reuse_rate"], 1), r(rec["water_intensity"], 2)])

# ---- 8. energy_data.xlsx ----------------------------------------------------
energy_headers = ["business_id", "month", "electricity_kwh", "grid_kwh", "renewable_kwh",
                  "diesel_litres", "LPG_kg", "peak_demand_kw", "production_units",
                  "energy_intensity_kwh_per_unit", "renewable_share_percent"]
energy_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        energy_rows.append([
            b["bid"], rec["month"], r(rec["electricity"]), r(rec["grid"]), r(rec["renewable"]),
            r(rec["genset_diesel"] + rec["process_diesel"]), r(rec["lpg"]), r(rec["peak_demand"]),
            r(rec["production"]), r(rec["energy_intensity"], 3), r(rec["renewable_share"], 1)])

# ---- 9. operations_materials.xlsx -------------------------------------------
ops_headers = ["business_id", "month", "raw_material_consumption_kg", "recycled_material_kg",
               "virgin_material_kg", "production_output_kg", "rejected_material_kg",
               "packaging_material_kg", "recycled_input_percent", "material_efficiency_percent"]
ops_rows = []
for b in BUSINESSES:
    for rec in DATA[b["bid"]]["rows"]:
        if rec["month"] not in LAST_12:
            continue
        ops_rows.append([
            b["bid"], rec["month"], r(rec["raw_material"]), r(rec["recycled_material"]),
            r(rec["virgin_material"]), r(rec["output_kg"]), r(rec["rejected"]),
            r(rec["packaging"]), r(rec["recycled_input_pct"], 1),
            r(rec["material_efficiency"], 1)])

# ---- 10. green_solutions.xlsx ----------------------------------------------
sol_headers = ["solution_id", "solution_name", "category", "description", "applicable_industries",
               "estimated_capex", "estimated_capex_min", "estimated_capex_max",
               "estimated_annual_savings",
               "estimated_annual_emission_reduction_tco2e", "estimated_water_saving_litres",
               "estimated_waste_reduction_kg", "estimated_annual_energy_saving_kwh",
               "estimated_payback_months", "implementation_months", "complexity",
               "expected_life_years", "in_climacred_catalog", "basis_note"]
sol_rows = []
for sid, s in SOLUTIONS.items():
    sol_rows.append([
        s["solution_id"], s["solution_name"], s["category"], s["description"],
        s["applicable_industries"], s["capex_mid"], s["capex_min"], s["capex_max"], s["savings"],
        r(s["co2"], 2), s["water_l"], s["waste_kg"], r(s["kwh_saving"]),
        r(s["payback"], 1), s["impl_months"], s["complexity"], s["life"], s["in_app"],
        ("Illustrative estimate for testing. Payback = mid-point capex / annual savings x 12. "
         "Savings and reductions are indicative, not guaranteed. " + s["assumptions"])])

# ---- 11. solution_recommendations.xlsx --------------------------------------
reco_headers = ["business_id", "solution_id", "solution_name", "priority", "reason",
                "baseline_metric", "estimated_improvement", "estimated_capex",
                "estimated_annual_savings", "estimated_emission_reduction",
                "estimated_payback_months", "implementation_timeline", "scaling_basis",
                "addresses_dimension", "addresses_dimension_score"]
reco_rows = []
for b in BUSINESSES:
    cur = DATA[b["bid"]]["rows"][-1]
    for rec in RECOMMENDATIONS[b["bid"]]:
        baseline = {
            "Energy": f"{cur['electricity']:,.0f} kWh/month; {cur['renewable_share']:.1f}% renewable",
            "Water": f"{cur['water']:,.0f} L/month; {cur['reuse_rate']:.1f}% reuse",
            "Waste": f"{cur['total_waste']:,.0f} kg/month; {cur['recycling_rate']:.1f}% recycled",
            "Mobility": f"{cur['fleet_diesel'] + cur['fleet_petrol']:,.0f} L fleet fuel/month",
            "Materials": f"{cur['raw_material']:,.0f} kg/month input; "
                         f"{cur['recycled_input_pct']:.1f}% recycled",
            "Operations": f"{b['facility_area_sqft']:,} sq ft; {b['hours']} h/day",
        }.get(rec["category"], "-")
        reco_rows.append([
            b["bid"], rec["solution_id"], rec["name"], rec["priority"], rec["reason"], baseline,
            rec["improvement"], rec["capex"], rec["savings"], r(rec["co2"], 2),
            r(rec["payback_months"], 1), rec["timeline"],
            f"Catalog values scaled x{rec['scale']:.2f} from this business's "
            f"{rec['category'].lower()} consumption",
            rec["addresses_dimension"], rec["dimension_score"]])

# ---- 12. scenarios.xlsx -----------------------------------------------------
scen_headers = ["scenario_id", "business_id", "scenario_name", "solutions_included", "capex",
                "annual_operating_cost_change", "annual_energy_savings_kwh",
                "annual_water_savings_litres", "annual_waste_reduction_kg",
                "annual_emission_reduction_tco2e", "annual_cost_savings", "payback_years",
                "projected_5_year_savings", "implementation_months",
                "baseline_annual_electricity_kwh", "baseline_annual_water_litres",
                "baseline_annual_emissions_tco2e"]
scen_rows = [[s["scenario_id"], s["business_id"], s["scenario_name"], s["solutions"], s["capex"],
              s["om"], s["energy_kwh"], s["water_l"], s["waste_kg"], s["co2"], s["savings"],
              s["payback_years"], s["five_year"], s["impl_months"], s["baseline_kwh"],
              s["baseline_water"], s["baseline_co2"]] for s in SCENARIOS]

# ---- 13. transformation_plans.xlsx ------------------------------------------
plan_headers = ["business_id", "phase", "action", "solution_id", "start_month",
                "duration_months", "expected_cost", "expected_savings",
                "expected_emission_reduction", "expected_water_saving", "owner_role",
                "dependency", "status"]
plan_rows = [[p["business_id"], p["phase"], p["action"], p["solution_id"], p["start_month"],
              p["duration_months"], p["expected_cost"], p["expected_savings"],
              p["expected_emission_reduction"], p["expected_water_saving"], p["owner_role"],
              p["dependency"], p["status"]] for p in PLAN]

# ---- 14. before_after_impact.xlsx -------------------------------------------
impact_headers = ["business_id", "metric", "baseline_value", "after_value", "unit",
                  "improvement_percent", "verification_status", "measurement_period",
                  "direction"]
impact_out = [[i["business_id"], i["metric"], i["baseline_value"], i["after_value"], i["unit"],
               i["improvement_percent"], i["verification_status"], i["measurement_period"],
               i["direction"]] for i in IMPACT]

# ---- 15. historical_climate_data.xlsx (24 months) ---------------------------
hist_headers = ["business_id", "month", "climate_score", "electricity_kwh", "water_litres",
                "waste_kg", "emissions_tco2e", "renewable_share_percent",
                "recycling_rate_percent", "water_reuse_rate_percent"]
hist_rows = []
for b in BUSINESSES:
    for rec, fp in zip(DATA[b["bid"]]["rows"], DATA[b["bid"]]["fps"]):
        hist_rows.append([
            b["bid"], rec["month"], fp["overallScore"], r(rec["electricity"]), r(rec["water"]),
            r(rec["total_waste"]), r(rec["total_emissions"], 2), r(rec["renewable_share"], 1),
            r(rec["recycling_rate"], 1), r(rec["reuse_rate"], 1)])

# ---- 16. gemini_ai_test_questions.xlsx ---------------------------------------
Q = []


def q(group, question, basis, forbidden):
    Q.append((f"Q{len(Q)+1:03d}", group, question, basis, forbidden))


q("A. Business Overview", "Give me a snapshot of my business profile and current climate position.",
  "business_profiles row + climate_assessments overall_climate_score for the loaded business",
  "No invented employee count, area or turnover")
q("A. Business Overview", "Which dimension of my climate fingerprint is the weakest?",
  "climate_assessments dimension scores (lowest of the six)",
  "No dimension scores other than stored ones")
q("A. Business Overview", "How complete is my climate data?",
  "data completeness from the assessment / profile completeness check",
  "Do not claim data exists that was never submitted")
q("A. Business Overview", "What is my climate risk level and why?",
  "risk_level + key_issue_1..3 from climate_assessments",
  "No new risk categories")
q("A. Business Overview", "Explain my climate fingerprint.",
  "six dimension scores + impact levels from the stored fingerprint",
  "No scores not present in the fingerprint")
q("A. Business Overview", "How do I compare with other businesses in my industry?",
  "Only if a benchmark is stored; otherwise state no benchmark data is available",
  "Must NOT invent an industry average or percentile ranking")
q("B. Energy", "Why is my energy score low?",
  "energy_score + energy_data electricity_kwh / energy_intensity_kwh_per_unit",
  "No kWh figures outside energy_data")
q("B. Energy", "How much electricity did I use last month and how does that compare to last year?",
  "energy_data for the latest month and the same month 12 months earlier (historical_climate_data)",
  "Do not compute a comparison if history is missing")
q("B. Energy", "What is my renewable energy share?",
  "renewable_share_percent = renewable_kwh / electricity_kwh x 100",
  "No assumed solar capacity")
q("B. Energy", "What is my energy intensity per unit of production?",
  "energy_intensity_kwh_per_unit from energy_data",
  "No new production figures")
q("B. Energy", "How much am I spending on electricity each month?",
  "Only if an electricity cost is stored; otherwise say the value is unavailable",
  "Must NOT invent a tariff or a monthly bill")
q("B. Energy", "Which months of the year use the most energy?",
  "historical_climate_data / energy_data monthly seasonality",
  "No months outside the stored range")
q("C. Water", "How much water am I consuming?",
  "water_data freshwater_intake_litres for the latest month",
  "No intake figure outside water_data")
q("C. Water", "What is my water reuse rate and how is it calculated?",
  "reused_water_litres / (freshwater_intake + reused) x 100",
  "No reuse rate not in the table")
q("C. Water", "How much wastewater do I generate and how much is treated?",
  "wastewater_generated_litres and wastewater_treated_litres",
  "No treatment efficiency figure not derivable from the two columns")
q("C. Water", "What is my water intensity per production unit?",
  "water_intensity_litres_per_unit",
  "No new production numbers")
q("C. Water", "Is my water use seasonal?",
  "12-24 month freshwater_intake_litres series",
  "Do not describe seasonality when fewer than 12 months are stored")
q("D. Waste", "Show me the main waste problem.",
  "waste_data streams: organic_kg / hazardous_kg / recyclable_kg vs recycling_rate_percent",
  "No stream not stored")
q("D. Waste", "What is my recycling rate and how is it computed?",
  "recycling_rate_percent and its inputs (recycled_kg, total_waste_kg)",
  "No rate outside the table")
q("D. Waste", "How much waste goes to landfill each month?",
  "landfill_kg from waste_data",
  "No landfill tonnage not stored")
q("D. Waste", "How much does waste disposal cost me?",
  "disposal_cost from waste_data",
  "No invented disposal tariff")
q("D. Waste", "How much hazardous waste do I generate and is it segregated?",
  "hazardous_kg plus the stored segregation flag",
  "Must not claim compliance or authorised-disposal status")
q("E. Emissions", "What are my biggest sources of emissions?",
  "emissions_data components: diesel_emissions, electricity_emissions, scope_3",
  "No source not represented in the data")
q("E. Emissions", "What is my total monthly footprint in tCO2e?",
  "total_emissions_tco2e = scope_1 + scope_2 + scope_3",
  "No total that is not the sum of the stored scopes")
q("E. Emissions", "What is my emissions intensity?",
  "emissions_intensity (kg CO2e per production unit)",
  "No new intensity calculation basis")
q("E. Emissions", "What share of my footprint is Scope 3?",
  "scope_3_emissions_tco2e / total_emissions_tco2e",
  "No upstream factors not documented")
q("E. Emissions", "Have my emissions gone up or down this year?",
  "historical_climate_data emissions_tco2e across 24 months",
  "Do not answer without stored history")
q("F. Mobility", "How can I reduce diesel consumption?",
  "mobility_data diesel_litres + recommendations for route optimization / EV transition",
  "No savings figures outside the recommendation data")
q("F. Mobility", "What are my mobility emissions?",
  "estimated_mobility_emissions_tco2e from mobility_data",
  "No fleet size not stored")
q("F. Mobility", "Is an EV fleet realistic for my operation?",
  "company_vehicle_count, diesel_litres, logistics_km + sol-ev-fleet economics",
  "Must label payback as an estimate, never a guarantee")
q("F. Mobility", "How much do employee commuting and business travel add?",
  "employee_commute_km and business_travel_km from mobility_data",
  "No commute distances not stored")
q("G. Recommendations", "What is my biggest climate inefficiency?",
  "worst dimension in climate_assessments + the matching baseline metric",
  "No inefficiency not supported by stored metrics")
q("G. Recommendations", "Which intervention should I investigate first and why?",
  "solution_recommendations priority 1 with its reason and baseline_metric",
  "No recommendation outside the stored recommendation set")
q("G. Recommendations", "Which solution has the shortest estimated payback?",
  "minimum estimated_payback_months in solution_recommendations",
  "Payback must be quoted as an estimate")
q("G. Recommendations", "What can I do with a budget of 10 lakh rupees?",
  "recommendations whose estimated_capex fits the stated budget",
  "Must not add solutions outside the catalog")
q("G. Recommendations", "Are these savings guaranteed?",
  "Explicit statement that savings are illustrative estimates, not guarantees",
  "MUST NOT say guaranteed, assured or certain savings")
q("G. Recommendations", "Can I claim carbon credits or certification from these actions?",
  "Explicit statement that no certification or carbon credit is issued by ClimaCred",
  "MUST NOT claim certification, verified credits or registry issuance")
q("H. Scenario Simulation", "Compare the solar and energy-efficiency scenarios.",
  "scenarios rows for the loaded business (capex, savings, payback, emission reduction)",
  "No scenario values outside scenarios.xlsx")
q("H. Scenario Simulation", "What happens if I install rooftop solar?",
  "Solar Scenario row: capex, annual_cost_savings, annual_emission_reduction, payback_years",
  "No generation figure outside the scenario data")
q("H. Scenario Simulation", "What does the combined transformation cost and return?",
  "Combined Transformation Scenario row",
  "No ROI or IRR that is not computed from stored columns")
q("H. Scenario Simulation", "Which scenario reduces the most emissions per rupee invested?",
  "annual_emission_reduction / capex across the stored scenarios",
  "Only compare scenarios that exist in the table")
q("I. Transformation Plan", "Create a 12-month transformation roadmap.",
  "transformation_plans rows ordered by phase and start_month",
  "No actions outside the stored plan")
q("I. Transformation Plan", "What should I complete in the first quarter?",
  "Phase 1 rows with start_month and duration_months",
  "No new timeline")
q("I. Transformation Plan", "Which plan actions depend on something else finishing first?",
  "dependency column of transformation_plans",
  "No invented dependencies")
q("I. Transformation Plan", "What is the total cost of my full plan?",
  "sum of expected_cost in transformation_plans for the business",
  "No capex outside the plan")
q("J. Historical Trends", "Has my climate performance improved over the last 12 months?",
  "historical_climate_data climate_score first vs last month of the window",
  "Do not answer with fewer than 12 stored months")
q("J. Historical Trends", "What changed compared with last month?",
  "latest vs previous month rows in historical_climate_data",
  "No month not stored")
q("J. Historical Trends", "Predict my electricity use for the next three months.",
  "Forecast only from stored monthly history; must state it is a trend projection, not a forecast of record",
  "Must not present a forecast as a measured or guaranteed value")
q("J. Historical Trends", "Why did my score change in the last six months?",
  "climate_score series plus the underlying resource columns",
  "No cause not supported by stored data")
q("K. Impact Verification", "Did my interventions actually reduce consumption?",
  "before_after_impact baseline_value vs after_value and improvement_percent",
  "Must use 'reported/observed change', never 'verified carbon reduction'")
q("K. Impact Verification", "How much water did the interventions save?",
  "before_after_impact 'Water consumption' row",
  "No saving outside the verification table")
q("K. Impact Verification", "Is this improvement verified by a third party?",
  "verification_status column - self-reported, not third-party verified",
  "MUST NOT claim third-party verification or assurance")
q("L. General Climate", "What is Scope 1, 2 and 3 in simple terms?",
  "General explanation; no business numbers needed",
  "No business numbers may be introduced")
q("L. General Climate", "What emission factors does ClimaCred use?",
  "Application configuration factors (electricity, diesel, petrol, gas)",
  "Must not present factors as officially certified values")
q("L. General Climate", "How is my climate score calculated?",
  "Documented dimension weights and 0-100 scoring range",
  "No claim of certification")
q("L. General Climate", "Which Indian subsidy scheme should I apply for?",
  "Out of scope - must state that scheme eligibility is not verified by ClimaCred",
  "MUST NOT assert eligibility, approval or financing")
q("L. General Climate", "Can you get my project financed or approved?",
  "Out of scope - must state that ClimaCred does not approve or arrange financing",
  "MUST NOT promise financing approval")

question_headers = ["question_id", "group", "question", "expected_answer_basis",
                    "must_not_contain", "applicable_businesses", "gemini_note"]
question_rows = []
for qid, group, question, basis, forbidden in Q:
    question_rows.append([
        qid, group, question, basis, forbidden,
        "Any loaded business (answer must use that business's stored data)",
        "Answer must cite only values ClimaCred supplies in context; where a value is absent the "
        "assistant must say the data is unavailable instead of estimating."])

# ---- 17. gemini_expected_behaviors.xlsx --------------------------------------
behavior_headers = ["test_id", "scenario", "user_question", "expected_behavior", "data_source",
                    "should_use_gemini", "should_refuse_hallucination", "expected_context",
                    "pass_condition"]
B = []


def beh(scenario, question, expected, source, use_gemini, refuse, context, pass_cond):
    B.append((f"AI-{len(B)+1:02d}", scenario, question, expected, source, use_gemini, refuse,
              context, pass_cond))


beh("No business profile exists",
    "What is my biggest climate inefficiency?",
    "State clearly that no business climate data is available and ask the user to complete the "
    "Business Profile and Climate Assessment. No business, industry or metric may be described.",
    "GET /api/profile -> null; GET /api/assessment -> null", "Yes", "Yes",
    "data_state.has_data = false",
    "Answer contains no numbers and no assumed industry; UI shows the empty state, not demo data.")
beh("Profile exists but no assessment",
    "Why is my energy score low?",
    "Explain that the Climate Assessment is incomplete, so no energy score exists yet.",
    "GET /api/profile (populated); GET /api/fingerprint -> null", "Yes", "Yes",
    "Profile present, assessment empty",
    "No energy score or kWh figure is invented; the response asks for the assessment.")
beh("Business has high energy consumption (B001 or B005)",
    "Why is my energy score low?",
    "Explain the stored energy score using the stored monthly kWh, kWh per production unit and "
    "renewable share - all quoted from the data.",
    "climate_assessments.energy_score + energy_data (electricity_kwh, energy_intensity, "
    "renewable_share_percent)", "Yes", "Yes",
    "Energy score < 40 for B001/B005; > 65 for B004",
    "Every number in the answer appears in the supplied context; energy_score matches the stored "
    "value for the loaded business.")
beh("User asks for an unsupported metric",
    "What is my monthly natural gas bill in rupees?",
    "State that the application does not currently have that value and offer the metrics it does "
    "hold.",
    "No gas cost column exists in the pack", "Yes", "Yes",
    "Field absent from all tables",
    "Answer says the value is unavailable and invents no rupee amount.")
beh("User asks for the latest trend",
    "What changed compared with last month?",
    "Compare the two most recent stored monthly records and name the changed metrics.",
    "historical_climate_data (last two rows for the loaded business)", "Yes", "Yes",
    "At least 24 monthly rows for the loaded business",
    "Both months quoted exist in historical_climate_data and the direction of change matches.")
beh("User asks for a recommendation",
    "Which intervention should I investigate first and why?",
    "Return the priority-1 stored recommendation with its reason, capex, savings and payback, and "
    "label them as estimates.",
    "solution_recommendations (priority = 1) + green_solutions", "Yes", "Yes",
    "At least 3 recommendation rows for the loaded business",
    "solution_id, capex, savings and payback match the stored priority-1 row exactly.")
beh("User asks whether savings are guaranteed",
    "Are these savings guaranteed?",
    "Explain that all savings, payback and emission reductions are illustrative estimates from the "
    "solution catalog, not guarantees or commitments.",
    "green_solutions basis_note + recommendation disclaimer", "Yes", "Yes",
    "Catalog rows carry the illustrative-estimate note",
    "Answer contains the words estimate/indicative and never 'guaranteed', 'assured' or 'certain'.")
beh("User asks about carbon credits / certification",
    "Can you certify these reductions or issue carbon credits?",
    "State that ClimaCred does not certify reductions, issue credits or register projects, and "
    "point to the reported/observed terminology instead.",
    "before_after_impact.verification_status; report disclaimer", "Yes", "Yes",
    "verification_status = self-reported (not third-party verified)",
    "No claim of certification, verification, registry listing or credit issuance.")
beh("Gemini API unavailable",
    "What is my biggest climate risk?",
    "The AI layer reports the disconnected/error state and returns no answer text; the UI shows the "
    "Gemini error state instead of pretending Gemini answered.",
    "GET /api/ai/status -> status not 'connected'; POST /api/ai/chat -> answer null with a status",
    "No", "Yes", "GEMINI_API_KEY empty or invalid in backend/.env",
    "Chat shows the error/unavailable message, no fabricated answer is displayed, and the navbar "
    "indicator is not 'connected'.")
beh("Empty database",
    "Explain my climate fingerprint.",
    "Dashboard, fingerprint, recommendations, scenarios, plan, impact and report all render their "
    "empty states; the chat states that no business data exists.",
    "All collections empty (POST /api/profile/reset)", "Yes", "Yes",
    "No profile, assessment, fingerprint, scenario, plan, impact or report documents",
    "No numeric card, chart series or recommendation is rendered; no demo business appears.")
beh("Only 6 months of history loaded",
    "Has my climate performance improved over the last 12 months?",
    "Answer only over the months that exist and state that a full 12-month comparison is not "
    "available.",
    "historical_climate_data (partial load)", "Yes", "Yes",
    "Fewer than 12 monthly rows for the loaded business",
    "The answer names the actual number of months available and invents no missing months.")
beh("Forecast with insufficient history",
    "Predict my electricity use for the next three months.",
    "Explain that forecasting needs at least 8 stored observations (the engine's minimum) and "
    "describe the stored trend instead of a fake forecast.",
    "backend forecasting minimum (MIN_HISTORY_FOR_FORECAST = 8)", "Yes", "Yes",
    "Load fewer than 8 monthly rows",
    "No numeric forecast is produced; the limitation is stated.")
beh("Impact verification wording",
    "Is this improvement verified by a third party?",
    "Report the observed before/after change and state clearly that it is self-reported, not "
    "third-party verified.",
    "before_after_impact.verification_status + improvement_percent", "Yes", "Yes",
    "3 businesses carry before/after rows",
    "The phrase 'verified carbon reduction' is never used; improvement_percent matches the table.")
beh("Cross-file consistency probe",
    "Do my energy, emissions and history tables agree?",
    "Quote the stored values and note any documented modelling difference instead of silently "
    "recomputing.",
    "energy_data + emissions_data + historical_climate_data for the same month", "Yes", "Yes",
    "Same month present in all three tables",
    "Quoted electricity and emissions match the tables; the answer does not invent a reconciling "
    "figure.")
beh("Recommendation not applicable to this business",
    "Should I install a closed-loop water recycling plant?",
    "Answer from the stored recommendations for the loaded business; if it is not recommended, say "
    "so and cite the business's actual water figures.",
    "solution_recommendations + water_data for the loaded business", "Yes", "Yes",
    "B004 has low water use and no water-recycling recommendation",
    "The answer reflects whether sol-water-ro is in that business's recommendation list.")
beh("Scenario values must not be re-derived",
    "What does the combined transformation cost and return?",
    "Quote the stored Combined Transformation Scenario row (capex, savings, payback) rather than "
    "re-adding solution prices.",
    "scenarios.xlsx Combined Transformation Scenario row", "Yes", "Yes",
    "35 scenario rows (7 per business)",
    "capex, annual_cost_savings and payback_years equal the stored row values.")

# ---------------------------------------------------------------------------
# Write the pack
# ---------------------------------------------------------------------------
OUT_XLSX = OUT_DIR                      # workbooks live at the pack root
OUT_CSV = os.path.join(OUT_DIR, "csv")
OUT_JSON = os.path.join(OUT_DIR, "api_payloads")
for d in (OUT_CSV, OUT_JSON):
    os.makedirs(d, exist_ok=True)

TABLES = {
    "business_profiles": (profile_headers, profile_rows,
                          [("api_payload_sheet", ["business_id", "post_api_profile_json"], profile_api)]),
    "climate_assessments": (assess_headers, assess_rows,
                            [("assessment_api_payloads",
                              ["business_id", "post_api_assessment_json",
                               "app_computed_scope1_plus_scope2_tco2e_month",
                               "pack_total_emissions_tco2e_month",
                               "pack_scope1_plus_scope2_tco2e_month", "overall_climate_score"],
                              assess_api)]),
    "resource_consumption": (resource_headers, resource_rows, None),
    "emissions_data": (emissions_headers, emissions_rows, None),
    "mobility_data": (mobility_headers, mobility_rows, None),
    "waste_data": (waste_headers, waste_rows, None),
    "water_data": (water_headers, water_rows, None),
    "energy_data": (energy_headers, energy_rows, None),
    "operations_materials": (ops_headers, ops_rows, None),
    "green_solutions": (sol_headers, sol_rows, None),
    "solution_recommendations": (reco_headers, reco_rows, None),
    "scenarios": (scen_headers, scen_rows, None),
    "transformation_plans": (plan_headers, plan_rows, None),
    "before_after_impact": (impact_headers, impact_out, None),
    "historical_climate_data": (hist_headers, hist_rows, None),
    "gemini_ai_test_questions": (question_headers, question_rows, None),
    "gemini_expected_behaviors": (behavior_headers,
                                  [list(x) for x in B], None),
}

for name, (headers, rows, extra) in TABLES.items():
    sheets = {name: (headers, rows)}
    if extra:
        for sh_name, sh_headers, sh_rows in extra:
            sheets[sh_name] = (sh_headers, sh_rows)
    write_xlsx(os.path.join(OUT_XLSX, name + ".xlsx"), sheets)
    write_csv(os.path.join(OUT_CSV, name + ".csv"), headers, rows)
    if extra:
        for sh_name, sh_headers, sh_rows in extra:
            write_csv(os.path.join(OUT_CSV, f"{name}__{sh_name}.csv"), sh_headers, sh_rows)

# ---- Mongo-ready fingerprint history (24 snapshots per business) -----------
# Shape exactly matches what backend/app/services/fingerprint_service.py reads
# back for GET /api/climate-fingerprint/history, so the historical charts and
# trend detection can be tested with the full 24-month series.
CONSISTENCY = {}
for b in BUSINESSES:
    docs = []
    for rec, fp in zip(DATA[b["bid"]]["rows"], DATA[b["bid"]]["fps"]):
        doc = dict(fp)
        doc["user_id"] = "default"
        doc["business_id"] = b["bid"]
        doc["data_origin"] = "climacred_test_data_pack"
        # Timestamps come from the modelled month, not from the clock, so the whole
        # pack rebuilds byte-for-byte identically (the application stamps its own
        # real time when it regenerates a fingerprint).
        stamp = f"{rec['month']}-28T06:00:00Z"
        doc["created_at"] = {"$date": stamp}
        doc["updated_at"] = {"$date": stamp}
        doc["generated_at"] = stamp
        doc["source_month"] = rec["month"]
        docs.append(doc)
    with open(os.path.join(OUT_JSON, f"{b['bid']}_mongo_fingerprint_history.json"), "w",
              encoding="utf-8") as fh:
        json.dump({"collection": "climate_fingerprints", "user_id": "default",
                   "note": "Insert with insertMany (see README_TESTING_GUIDE.md section 6). "
                           "Fictional demo data for testing only.",
                   "documents": docs}, fh, indent=2, ensure_ascii=False, default=str)

    cur = DATA[b["bid"]]["rows"][-1]
    app_em = calculate_emissions_metrics(DATA[b["bid"]]["assessment"])
    CONSISTENCY[b["bid"]] = {
        "business_name": b["name"],
        "app_scope1_scope2_tco2e_month": app_em["emissions_breakdown_tonnes_co2e_per_month"]["total"],
        "pack_scope1_scope2_tco2e_month": round(cur["scope1"] + cur["scope2"], 3),
        "difference_tco2e_month": round(
            app_em["emissions_breakdown_tonnes_co2e_per_month"]["total"]
            - (cur["scope1"] + cur["scope2"]), 3),
        "has_solar": cur["solar_kw"] > 0,
        "solar_kw": cur["solar_kw"],
        "app_total_kwh": DATA[b["bid"]]["assessment"]["energy"]["monthlyElectricityKwh"],
        "grid_kwh": cur["grid"],
        "ev_kwh": cur["ev_kwh"],
        "pack_total_emissions_tco2e_month": round(cur["total_emissions"], 3),
        "app_uses_total_kwh_times_factor": True,
        "recycling_rate": round(cur["recycling_rate"], 2),
        "reuse_rate": round(cur["reuse_rate"], 2),
        "renewable_share": round(cur["renewable_share"], 2),
        "energy_score": DATA[b["bid"]]["fingerprint"]["dimensions"][0]["score"],
        "overall_score": DATA[b["bid"]]["fingerprint"]["overallScore"],
        "score_24m_first": DATA[b["bid"]]["fps"][0]["overallScore"],
        "score_24m_last": DATA[b["bid"]]["fps"][-1]["overallScore"],
        "solar_first_share": round(DATA[b["bid"]]["rows"][0]["renewable_share"], 1),
    }
with open(os.path.join(OUT_DIR, "generator", "_consistency_notes.json"), "w",
          encoding="utf-8") as fh:
    json.dump(CONSISTENCY, fh, indent=2)

# API payloads (one POST-ready file per business)
for b in BUSINESSES:
    payload = {
        "business_id": b["bid"],
        "note": "POST /api/profile then /api/assessment for this business, then "
                "POST /api/fingerprint/generate. Fictional demo data.",
        "profile": DATA[b["bid"]]["profile"],
        "assessment": DATA[b["bid"]]["assessment"],
        "expected_overall_climate_score": DATA[b["bid"]]["fingerprint"]["overallScore"],
    }
    with open(os.path.join(OUT_JSON, f"{b['bid']}_payload.json"), "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)

# ---------------------------------------------------------------------------
# MANIFEST
# ---------------------------------------------------------------------------
manifest_headers = ["file_name", "purpose", "entity", "row_count", "required_for_import",
                    "dependency", "test_scenario"]
manifest = [
    ["business_profiles.xlsx", "5 fictional Indian MSMEs (identity, size, location, contact "
     "placeholder). Import first.", "Business profile", len(profile_rows), "Yes", "None",
     "Profile page, dashboard header, empty-vs-loaded state"],
    ["climate_assessments.xlsx", "Current-month assessment scores per business, computed by the "
     "repository's own scoring engine; second sheet holds the POST-ready assessment JSON.",
     "Climate assessment / fingerprint", len(assess_rows), "Yes", "business_profiles.xlsx",
     "Climate Assessment page, Climate Fingerprint, dashboard scores"],
    ["resource_consumption.xlsx", "12 months of consolidated resource data per business (60 rows).",
     "Monthly resource ledger", len(resource_rows), "No",
     "business_profiles.xlsx", "Dashboard trend cards, monthly totals, data-quality checks"],
    ["emissions_data.xlsx", "Monthly Scope 1/2/3 with total = scope1+scope2+scope3.",
     "Emissions", len(emissions_rows), "No", "resource_consumption.xlsx",
     "Emissions page, source breakdown, intensity charts"],
    ["mobility_data.xlsx", "Monthly fleet fuels, EV kWh, commute/travel/logistics km.",
     "Mobility", len(mobility_rows), "No", "resource_consumption.xlsx",
     "Mobility page, diesel-heavy vs EV-adopting profiles"],
    ["waste_data.xlsx", "Monthly waste streams, recycled volume, landfill and disposal cost.",
     "Waste", len(waste_rows), "No", "resource_consumption.xlsx",
     "Waste page, recycling-rate math, circularity tests"],
    ["water_data.xlsx", "Monthly intake split, reuse, wastewater treated, water intensity.",
     "Water", len(water_rows), "No", "resource_consumption.xlsx",
     "Water page, reuse-rate math, seasonal variation"],
    ["energy_data.xlsx", "Monthly grid vs renewable split, peak demand, energy intensity.",
     "Energy", len(energy_rows), "No", "resource_consumption.xlsx",
     "Energy page, renewable share, inefficiency vs efficient business"],
    ["operations_materials.xlsx", "Monthly material input/output, recycled input, efficiency.",
     "Operations & materials", len(ops_rows), "No", "resource_consumption.xlsx",
     "Operations page, material-efficiency tests"],
    ["green_solutions.xlsx", "18-solution catalog: the 12 in-app solutions (exact ids) plus 6 "
     "extended illustrative solutions.", "Solution catalog", len(sol_rows), "No", "None",
     "Green Solutions page, recommendation economics, payback ordering"],
    ["solution_recommendations.xlsx", "3-6 recommendations per business derived from that "
     "business's own inefficiencies.", "Recommendations", len(reco_rows), "No",
     "climate_assessments.xlsx + green_solutions.xlsx",
     "Recommendation engine relevance, payback sorting, budget filtering"],
    ["scenarios.xlsx", "7 scenarios per business (35 rows) with internally consistent sums.",
     "Scenario simulation", len(scen_rows), "No", "solution_recommendations.xlsx",
     "Scenario Simulator, scenario comparison, 5-year projection"],
    ["transformation_plans.xlsx", "5-phase roadmap per business tied to solution ids.",
     "Transformation plan", len(plan_rows), "No", "solution_recommendations.xlsx",
     "Transformation Plan page, phase sequencing, dependencies, status"],
    ["before_after_impact.xlsx", "Baseline vs post-intervention metrics for B001, B003, B004.",
     "Impact verification", len(impact_out), "No", "historical_climate_data.xlsx",
     "Impact Verification page, improvement math, unverified-claim wording"],
    ["historical_climate_data.xlsx", "24 months per business (120 rows) for charts, trend "
     "detection and Gemini predictions.", "Historical series", len(hist_rows), "No",
     "business_profiles.xlsx", "Historical charts, trend detection, forecasting, fingerprint "
     "evolution, 'what changed'"],
    ["gemini_ai_test_questions.xlsx", f"{len(question_rows)} chat questions in 12 groups with the "
     "data each answer must come from.", "AI test questions", len(question_rows), "No", "All data files",
     "Gemini chat grounding, refusal of invented values"],
    ["gemini_expected_behaviors.xlsx", f"{len(B)} integration test cases including no-data, "
     "Gemini-disconnected and empty-database states.", "AI behaviour tests", len(B), "No",
     "All data files", "Gemini status states, hallucination guards, empty-state behaviour"],
    ["README_TESTING_GUIDE.md", "Import order, dependency map, 15 test procedures, factor table "
     "and validation summary.", "Documentation", 1, "Yes", "All files",
     "End-to-end test execution"],
]
write_xlsx(os.path.join(OUT_DIR, "MANIFEST.xlsx"), {"MANIFEST": (manifest_headers, manifest)})
write_csv(os.path.join(OUT_DIR, "MANIFEST.csv"), manifest_headers, manifest)

print("tables written:", len(TABLES))
print("rows:", {k: len(v[1]) for k, v in TABLES.items()})
print("manifest rows:", len(manifest))
for b in BUSINESSES:
    fp = DATA[b["bid"]]["fingerprint"]
    first = DATA[b["bid"]]["fps"][0]
    print(f"{b['bid']} score {first['overallScore']} -> {fp['overallScore']} "
          f"({fp['scoreLabel']}) recs={len(RECOMMENDATIONS[b['bid']])}")
